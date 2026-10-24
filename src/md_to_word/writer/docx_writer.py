"""md_to_word.writer.docx_writer — 文档模型 → .docx。

这是自研写出器：只做机械映射，不做任何判断——样式名、段落顺序、表格线宽
这些决定在读取阶段和格式契约里已经定完。写出一层如果再"猜"，就等于把判据藏到了渲染里。

能力边界写在 `SUPPORTED` 里，遇到没实现的结构抛 `UnsupportedFeature` 并说清楚缺什么，
**不做静默降级**：悄悄写成一个普通段落，比直接失败更难查。
"""

from __future__ import annotations

import copy
import re
from pathlib import Path
from typing import Iterable, Optional

from docx import Document as open_docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Emu

from md_to_word.config import default_crossref_prefixes, looks_chinese
from md_to_word.i18n import t
from md_to_word.utils.docx_zip import normalize_zip
from md_to_word.utils.pdf_figure import PdfFigureError, is_pdf, rasterize
from md_to_word.preprocessor.markers import MarkerError, section_paragraph_xml, toc_xml
from md_to_word.spec.constants import PAGE_NUMBER_FMT
from md_to_word.transform.mathml2omml import UnsupportedMath, latex_to_omml
from md_to_word.model.document import (
    Block, Code, CodeBlock, Document, Emphasis, Equation, Figure, FootnoteRef, Heading,
    Inline, LineBreak, Link, ListBlock, MAX_LIST_LEVEL, Math, PageBreak, Paragraph, Quote,
    SectionBreak, Strikethrough,
    Strong, Subscript, Superscript, Table, Text, Toc, Underline, Citation, CrossRef,
)

SUPPORTED_BLOCKS = (Heading, Paragraph, ListBlock, Table, Figure, CodeBlock, PageBreak, Quote)

# 列表缩进（twips，1 英寸 = 1440）：每层左缩进 0.5 英寸，编号悬挂 0.25 英寸。
# 引用块里的列表整体再右移一个正文缩进，免得和引用正文左对齐、看不出从属关系。
_LIST_STEP_TWIPS = 720
_LIST_HANGING_TWIPS = 360
_QUOTE_LEFT_TWIPS = 480
SUPPORTED_INLINES = (Code, Strong, Emphasis, Strikethrough, Underline, Superscript, Subscript)

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}


_W_FMT = {"bold": "b", "italic": "i", "strike": "strike", "underline": "u"}


def _plain_text(nodes: Iterable[Inline]) -> str:
    """行内节点里真正的文字（引用、交叉引用这类不是正文，判语种时不看它们）。"""
    out: list[str] = []
    for node in nodes:
        if isinstance(node, Text):
            out.append(node.value)
        else:
            children = getattr(node, "children", None)
            if children:
                out.append(_plain_text(children))
    return "".join(out)


def _ends_with_cjk(text: str) -> bool:
    """行尾是中文——汉字、或中文句末引号/括号都算（`满意” [13]` 里那个空格同样是语法）。"""
    stripped = text.rstrip()
    if not stripped:
        return False
    last = stripped[-1]
    return "\u4e00" <= last <= "\u9fff" or last in "”’」』）】》"


def _default_eq_format(chinese: bool, by_chapter: bool) -> str:
    """没写「公式编号格式」时的默认外观。

    中文稿全角括号、章内用短横（`（5-1）`，学位论文与多数中文期刊的写法）；
    英文稿半角括号、点号（`(5.1)`）。
    """
    if chinese:
        return "（{c}-{n}）" if by_chapter else "（{n}）"
    return "({c}.{n})" if by_chapter else "({n})"


def _escape(text: str) -> str:
    return (
        text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    )


def _trim_citation_gap(nodes: Iterable[Inline]) -> list:
    """中文正文与数字制引文之间的那个空格不该印出来：`认知负荷 [9]` → `认知负荷[9]`。

    稿子里的空格是 Markdown 语法要求（`@key` 前必须断开），不是作者想印的，因此不印。
    西文行文（`as shown in [9]`）照留——那里空格是本意。
    """
    out: list = []
    for node in nodes:
        if isinstance(node, Citation) and out and isinstance(out[-1], Text):
            prev = out[-1]
            if prev.value.endswith((" ", "\u00a0")):
                if not prev.value.strip() and len(out) >= 2:
                    if _ends_with_cjk(_plain_text([out[-2]])):
                        out.pop()
                elif _ends_with_cjk(prev.value):
                    out[-1] = Text(prev.value.rstrip(" \u00a0"))
        out.append(node)
    return out


class UnsupportedFeature(NotImplementedError):
    """模型里有、但写出器还没实现的结构。"""


class _Writer:
    def __init__(
        self,
        document: Document,
        docx,
        citations=None,
        base_dir=None,
        equation_number_placement="right",
        equation_number_format="",
        no_row_split: bool = False,
        repeat_header: bool = False,
        lang: Optional[str] = None,
    ) -> None:
        self.model = document
        self.docx = docx
        self.citations = citations
        self.equation_number_placement = equation_number_placement
        self.base_dir = Path(base_dir) if base_dir else Path.cwd()
        self._after_heading = False  # 下一段正文若是「标题后的第一段」，落 First Paragraph
        # 样式表在写出期间不变。python-docx 每次迭代 docx.styles 都会重建 StyleFactory，
        # 逐段按名查询占写出阶段约四成时间，故缓存名称、引用与对象。
        self._style_names: dict[str, str] | None = None
        self._style_refs: dict[str, str] | None = None
        self._style_objects: dict | None = None
        self._compact_style = self._style_name("Compact")
        self._footnote_ids: dict[str, int] = {}
        self._eq_numbers: dict[str, tuple[int, int]] = {}  # 标签 → (章号, 序号)
        self._last_paragraph = None  # 最近写入的段落，keepNext 要往上补
        self._bibliography_written = False  # 文献表只在「参考文献标题」后写一次
        self.no_row_split = no_row_split  # 表格行是否禁止跨页断开（要求文件写「行不跨页」）
        self.repeat_header = repeat_header  # 表头跨页重复（要求文件写「跨页重复」）
        chinese = looks_chinese(self._prose_text())
        # 报错语言：流水线按 `--lang`/稿件语种传进来；直接调库时按正文自己判
        self.lang = lang or ("zh" if chinese else "en")
        self._crossref_prefixes = {
            **default_crossref_prefixes(chinese),
            **(document.meta.crossref or {}),
        }
        # 编号外观：要求文件写了就照写的来，没写按语种给默认（中文全角短横、英文半角点）
        self.eq_number_format = equation_number_format or _default_eq_format(
            chinese, document.meta.chapters
        )

    # ── 入口 ─────────────────────────────────────────────────────────
    def write(self) -> None:
        body = self.docx.element.body
        for paragraph in list(self.docx.paragraphs):  # 模板自带的空段落
            if not paragraph.text.strip():
                body.remove(paragraph._element)

        self._collect_equation_numbers()
        blocks = self.model.blocks
        for index, block in enumerate(blocks):
            self.block(block)
            self._keep_caption_with_figure(blocks, index)
        self._write_bibliography()
        if self._footnote_ids:
            self._write_footnotes()

    # ── 块 ───────────────────────────────────────────────────────────
    def block(self, block: Block) -> None:
        """任何非段落块之后，下一段正文都算「首段」——这是 Word 里 First Paragraph 的用法：
        公式、图、表、列表之后的那一段，与标题之后一样换样式。"""
        prev = self._after_heading
        if isinstance(block, Heading):
            self._heading(block)
        elif isinstance(block, Paragraph):
            self._paragraph(block)
        elif isinstance(block, Table):
            self._table(block)
        elif isinstance(block, Figure):
            self._figure(block)
        elif isinstance(block, CodeBlock):
            self._code(block)
        elif isinstance(block, ListBlock):
            self._list(block)
        elif isinstance(block, Quote):
            self._quote(block)
        elif isinstance(block, PageBreak):
            self.docx.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif isinstance(block, Equation):
            self._equation(block)
        elif isinstance(block, Toc):
            try:
                toc = toc_xml(block.depth or self._toc_depth(), self.lang)
            except MarkerError as exc:
                # 级数写超（`{{< toc 99 >}}`）是稿件写法问题，不该以 traceback 收场
                raise UnsupportedFeature(t("writer.toc_failed", self.lang, err=exc)) from exc
            self._insert_before_sectpr(parse_xml(toc))
        elif isinstance(block, SectionBreak):
            self._section(block)
        else:
            raise UnsupportedFeature(t("writer.unsupported_block", self.lang,
                                       type=type(block).__name__))

        if not isinstance(block, Paragraph):
            self._after_heading = True
        elif prev:
            self._after_heading = False

    def _write_bibliography(self) -> None:
        """文献列表由引擎生成；样式由要求文件的「参考文献」行控制。

        插在 `%% 参考文献标题` 那一行**之后**，而不是无条件押到文末——文献后面接附录、
        接致谢都是正常的稿子结构，押到文末会把它们排到文献列表前面去。
        """
        if self._bibliography_written or self.citations is None:
            return
        self._bibliography_written = True
        entries = self.citations.bibliography_entries()
        if not entries:
            return
        for key, segments in entries:
            paragraph = self._new_paragraph(self._style_ref("Bibliography"))
            self._bookmark(paragraph, f"_Ref_{key}", segments)

    _bookmark_seq = 100

    def _bookmark(self, paragraph, name: str, segments) -> None:
        """在文献条目开头放书签，正文引用才能跳过来。条目里的斜体由 CSL 决定，要保留。"""
        type(self)._bookmark_seq += 1
        bid = type(self)._bookmark_seq
        paragraph._p.append(
            parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{bid}" w:name="{name}"/>')
        )
        for text, fmt in segments:
            paragraph._p.append(self._run_element(paragraph._p, text, fmt))
        paragraph._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{bid}"/>'))

    def _collect_equation_numbers(self) -> None:
        """先扫一遍给带标签的公式编号：写正文时要用它解析 `@eq-x`。

        编号规则：默认按出现顺序 (1)(2)…；开了 `chapters: true` 就按「章号, 序号」，
        章号取自一级标题的序号。外观由 `equation_number_format` 模板决定。
        """
        chapter = 0
        index = 0
        for block in self.model.blocks:
            if isinstance(block, Heading) and block.level == 1:
                chapter += 1
                index = 0
            elif isinstance(block, Equation) and block.label:
                index += 1
                self._eq_numbers[block.label] = (chapter, index)

    def _content_width_twips(self) -> int:
        sect = self._body_sectpr()
        if sect is None:
            return 8306
        size = sect.find(qn("w:pgSz"))
        margins = sect.find(qn("w:pgMar"))
        try:
            width = int(size.get(qn("w:w")))
            return width - int(margins.get(qn("w:left"))) - int(margins.get(qn("w:right")))
        except (AttributeError, TypeError, ValueError):
            return 8306

    def _body_sectpr(self):
        body = self.docx.element.body
        # body 的 sectPr 固定在末尾，先看末位可免遍历
        if len(body) and body[-1].tag == qn("w:sectPr"):
            return body[-1]
        return body.find(qn("w:sectPr"))

    def _insert_before_sectpr(self, element) -> None:
        """插到文末 sectPr 之前：body 的 sectPr 必须是最后一个子元素。"""
        body = self.docx.element.body
        sectpr = self._body_sectpr()
        if sectpr is not None:
            sectpr.addprevious(element)
        else:
            body.append(element)

    def _toc_depth(self) -> int:
        """目录收录到第几级：模板里声明了 目录一/二/三…级 就收到几级。

        声明与收录必须一致——要求文件写了「目录四级」，目录却只收到三级，
        用户会以为第四条样式没生效；反之声明没写到，多收的那级只能退回 Word 默认。
        模板里一条都没有时按 3 级（Word 的默认口径）。
        """
        names = {s.name.lower() for s in self.docx.styles}  # 内置样式名是小写 `toc 1`
        levels = [n for n in range(1, 6) if f"toc {n}" in names]
        return max(levels) if levels else 3

    def _section(self, block: SectionBreak) -> None:
        """分节符：它描述的是**上面那一节**（OOXML 里 sectPr 归前一节）。"""
        params: dict = {}
        fmt = PAGE_NUMBER_FMT.get(block.page_format)
        if fmt:
            params["fmt"] = fmt
        if block.start:
            params["start"] = block.start

        template_sect = self._body_sectpr()
        xml = section_paragraph_xml(params, copy.deepcopy(template_sect) if template_sect is not None else None)
        self._insert_before_sectpr(parse_xml(xml))

    def _equation(self, block: Equation) -> None:
        """独立公式：居中，编号靠右。编号用中心制表位 + 右制表位实现，是 Word 里的常规写法。

        没有编号时段落里**只放公式**：Word 按「段落里除公式外还有没有别的东西」决定
        按显示公式还是行内公式排。带编号的公式与制表位同段，只能按行内排出（分式缩小、
        求和的上下限挪到右侧）——这是 Word 的规矩，不是这里能绕开的。
        """
        try:
            maths = latex_to_omml(block.latex, self.lang)
        except UnsupportedMath as exc:
            snippet = block.latex.strip().replace("\n", " ")[:40]
            raise UnsupportedFeature(t("writer.equation_failed", self.lang,
                                       snippet=snippet, err=exc)) from exc

        paragraph = self._new_paragraph("Body Text")
        number = self._eq_numbers.get(block.label or "")
        if number is None:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph._p.append(self._omath(maths))
            return

        width = self._content_width_twips()
        pPr = paragraph._p.get_or_add_pPr()
        # 中心制表位永远要有（公式居中）；右制表位只在编号靠右时才需要
        right_tab = (
            "" if self.equation_number_placement == "inline"
            else f'<w:tab w:val="right" w:pos="{width}"/>'
        )
        pPr.append(
            parse_xml(
                f'<w:tabs {nsdecls("w")}><w:tab w:val="center" w:pos="{width // 2}"/>'
                f"{right_tab}</w:tabs>"
            )
        )
        paragraph.add_run().add_tab()
        math_el = self._omath(maths)
        paragraph._p.append(math_el)

        text = self._render_eq_number(*number)
        if self.equation_number_placement == "inline":
            paragraph.add_run(text)
        else:
            paragraph.add_run().add_tab()
            paragraph.add_run(text)

    @staticmethod
    def _omath(nodes):
        math_el = parse_xml(
            f'<m:oMath {nsdecls("m")} xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
        )
        for node in nodes:
            math_el.append(node)
        return math_el

    def _heading(self, block: Heading) -> None:
        paragraph = self._new_paragraph(f"Heading {block.level}")
        self.inlines(paragraph, block.inlines)
        self._after_heading = True
        clean_title = re.sub(
            r"^[0-9一二三四五六七八九十第\s.、章节附录]+", "", _plain_text(block.inlines).strip()
        ).strip()
        if clean_title in ("参考文献", "参考书目", "References", "Bibliography"):
            self._write_bibliography()

    def _paragraph(self, block: Paragraph) -> None:
        """未指定样式的正文，按契约分流：紧跟标题的第一段用 First Paragraph，其余 Body Text。"""
        if block.style:
            style = block.style
        else:
            style = "First Paragraph" if self._after_heading else "Body Text"
        self._after_heading = False
        # 样式在模板里不存在时退回 Normal：预检本该拦住，走到这里说明门禁被关掉了，
        # 此时宁可降级排版，也不要 KeyError 中断整份稿件的导出。
        paragraph = self._new_paragraph(self._style_name(style) or "Normal")
        self._last_paragraph = paragraph
        self.inlines(paragraph, block.inlines)
        if block.style == "References Title":
            self._write_bibliography()

    def _code(self, block: CodeBlock) -> None:
        paragraph = self._new_paragraph(self._style_ref("SourceCode"))
        lines = block.value.split("\n")
        for i, line in enumerate(lines):
            if i:
                paragraph.add_run().add_break()
            paragraph.add_run(line)

    def _quote(self, block: Quote) -> None:
        for inner in block.blocks:
            if isinstance(inner, Paragraph):
                paragraph = self._new_paragraph(inner.style or "Block Text")
                self.inlines(paragraph, inner.inlines)
            elif isinstance(inner, ListBlock):
                # 引用里的列表整体再缩进一个层级，免得和引用里的正文左对齐、分不出从属关系
                self._list(inner, base_left=_QUOTE_LEFT_TWIPS)
            elif isinstance(inner, CodeBlock):
                self._code(inner)
            else:
                raise UnsupportedFeature(t("writer.quote_children", self.lang))

    def _list(self, block: ListBlock, base_left: int = 0) -> None:
        """列表写出：真实编号（numbering.xml），不是字面量序号。

        字面量的代价是：Word 里按退格改不了级别、增删条目不会自动重编号、目录与交叉引用也
        抓不到。既然目标是「导出的文档还能在 Word 里继续改」，编号就得是域。
        """
        body_style = self._style_name("Body Text") or self._compact_style
        # 每一层各自的起始号（读取端按层记在块上）：只写 ilvl=0 的话，深层会从 1 重数，
        # 作者写的 `3.` 到 Word 里变成 a)。
        starts = getattr(block, "level_starts", None) or {0: block.start}
        # numId 必须整块列表共用一个：`_numbered_list` 遇到起始号不是 1 时会**新建**一个
        # 带 startOverride 的编号实例，逐项调用就会逐项重新从起始号数起——「4. 5.」变成「4. 4.」。
        num_ids: dict[bool, int] = {}
        for item in block.items:
            ordered = block.ordered if item.ordered is None else item.ordered
            if ordered not in num_ids:
                num_ids[ordered] = self._numbered_list(ordered, starts)
            num_id = num_ids[ordered]
            for inner in item.blocks:
                if not isinstance(inner, Paragraph):
                    raise UnsupportedFeature(t("writer.list_item_children", self.lang))
                # 列表项是正文，不是表格单元格：套 Body Text 而不是 Compact。
                # Compact 常被设成「居中 + 不缩进」（表格里就该这样），列表项跟着它
                # 会被逐行居中，看着像手打的「• 一句话」——排版对不上，内容却没错。
                paragraph = self._new_paragraph(body_style)
                self._apply_numbering(paragraph, num_id, item.level)
                # 缩进写在段落上：左缩进按层递增、悬挂 0.25 英寸，与编号定义一致；
                # 同时清掉正文样式继承来的首行缩进（含字符单位的 firstLineChars），
                # 否则编号会被推后两字。段落级 w:ind 会覆盖编号定义里的那一份，故两处都得写。
                ind = paragraph._p.get_or_add_pPr().get_or_add_ind()
                ind.set(qn("w:left"), str(base_left + _LIST_STEP_TWIPS * (item.level + 1)))
                ind.set(qn("w:hanging"), str(_LIST_HANGING_TWIPS))
                ind.set(qn("w:leftChars"), "0")
                ind.set(qn("w:firstLineChars"), "0")
                self.inlines(paragraph, inner.inlines)

    def _ensure_numbering(self) -> None:
        """模板里的 numbering.xml 是空的：首次用到列表时补两套定义（无序/有序），各三级。"""
        numbering = self.docx.part.numbering_part.element
        if numbering.findall(qn("w:abstractNum")):
            return
        bullet_marks = ("•", "–", "▪")  # 第二级用短横：空心圆在中西文混排里容易被看成一个字母
        ordered_fmt = (("decimal", "%1."), ("lowerLetter", "%2)"), ("lowerRoman", "%3."))
        levels = []
        for abstract_id, marks in ((0, bullet_marks), (1, None)):
            lvl_xml = []
            for ilvl in range(MAX_LIST_LEVEL + 1):
                indent = f'<w:pPr><w:ind w:left="{_LIST_STEP_TWIPS * (ilvl + 1)}" w:hanging="{_LIST_HANGING_TWIPS}"/></w:pPr>'
                if marks is None:
                    fmt, text = ordered_fmt[ilvl]
                else:
                    fmt, text = "bullet", marks[ilvl]
                lvl_xml.append(
                    f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
                    f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>{indent}</w:lvl>'
                )
            levels.append(
                f'<w:abstractNum w:abstractNumId="{abstract_id}">'
                f'<w:multiLevelType w:val="multilevel"/>{"".join(lvl_xml)}</w:abstractNum>'
            )
        xml = (
            f'<w:numbering {nsdecls("w")}>'
            + "".join(levels)
            + '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'
            '<w:num w:numId="2"><w:abstractNumId w:val="1"/></w:num>'
            "</w:numbering>"
        )
        for child in parse_xml(xml):
            numbering.append(child)

    def _numbered_list(self, ordered: bool, starts: dict[int, int]) -> int:
        """返回该列表要用的 numId。有序列表各自一份，编号才从 1（或指定值）重新开始。"""
        self._ensure_numbering()
        base_id = 2 if ordered else 1
        if not ordered:
            return base_id
        # 有序列表每个都要自己的 numId：共用一个的话，第二个列表会接着上一个往下编号
        # （Word 里就是 3.、4.），而且起始号也被忽略。

        # 起始号按层写：深层要拿到自己的 lvlOverride，只写 ilvl=0 只能救顶层。
        # 起始号本来就是 1 的层（抽象定义里的默认值）不写，numbering.xml 不因此变胖。
        levels = [0] + [lv for lv in sorted(starts) if lv > 0 and starts[lv] != 1]
        overrides = "".join(
            f'<w:lvlOverride w:ilvl="{lv}"><w:startOverride w:val="{starts.get(lv, 1)}"/></w:lvlOverride>'
            for lv in levels
        )
        numbering = self.docx.part.numbering_part.element
        self._override_seq = getattr(self, "_override_seq", 99)
        self._override_seq += 1
        new_id = self._override_seq
        xml = (
            f'<w:num {nsdecls("w")} w:numId="{new_id}">'
            f'<w:abstractNumId w:val="{base_id - 1}"/>'
            f"{overrides}"
            "</w:num>"
        )
        numbering.append(parse_xml(xml))
        return new_id

    @staticmethod
    def _apply_numbering(paragraph, num_id: int, level: int = 0) -> None:
        pPr = paragraph._p.get_or_add_pPr()
        numPr = pPr.get_or_add_numPr()
        numPr.get_or_add_ilvl().val = level
        numPr.get_or_add_numId().val = num_id

    def _table(self, block: Table) -> None:
        rows = len(block.rows) + (1 if block.header else 0)
        cols = len(block.header) if block.header else (len(block.rows[0]) if block.rows else 0)
        if not rows or not cols:
            return
        self._separate_from_previous_table()
        table = self.docx.add_table(rows=rows, cols=cols)
        table.style = self._style_object("Table")
        table.autofit = block.widths is None

        cell_style = self._style_name("Compact")
        data = ([block.header] if block.header else []) + block.rows
        for r, row in enumerate(data):
            for c, cell_inlines in enumerate(row):
                if c >= cols:
                    break
                cell = table.cell(r, c)
                paragraph = cell.paragraphs[0]
                if cell_style:
                    paragraph.style = self._style_object(cell_style)
                if block.alignments and c < len(block.alignments):
                    paragraph.alignment = _ALIGN.get(block.alignments[c], WD_ALIGN_PARAGRAPH.LEFT)
                self.inlines(paragraph, cell_inlines)
        if self.no_row_split:
            for row in table.rows:
                trPr = row._tr.get_or_add_trPr()
                if trPr.find(qn("w:cantSplit")) is None:
                    trPr.append(trPr.makeelement(qn("w:cantSplit"), {}))
        if block.header and self.repeat_header:
            self._mark_header_row(table)
        if block.widths:
            self._apply_column_widths(table, block.widths)

    def _separate_from_previous_table(self) -> None:
        """两张表紧挨着时垫一个空段落。

        OOXML 没有「两张相邻的表」：`</w:tbl><w:tbl>` 在 Word 里会被并成一张——第二张的
        表头、题注、独立编号全部跟着错位，而且不报错。中间隔一个 1 磅的空段落就分得开，
        高度可以忽略。
        """
        body = self.docx.element.body
        sectpr = self._body_sectpr()
        # 只看 sectPr 的前一个元素；遍历整个 body 的代价是「表数 × 段落数」
        last = sectpr.getprevious() if sectpr is not None else (body[-1] if len(body) else None)
        if last is None or last.tag != qn("w:tbl"):
            return
        self._insert_before_sectpr(
            parse_xml(
                f'<w:p {nsdecls("w")}><w:pPr><w:spacing w:before="0" w:after="0" '
                f'w:line="240" w:lineRule="auto"/><w:rPr><w:sz w:val="2"/></w:rPr></w:pPr></w:p>'
            )
        )

    def _apply_column_widths(self, table, widths: list[float]) -> None:
        """分隔行短横线长短 → 列宽比例，按版心铺开。

        `gridCol`、`tblW`、每格的 `tcW` 三处都要写：Word 以 `tcW` 为准、`gridCol` 决定
        网格，`tblLayout=fixed` 不钉住的话存盘时会被重算——少写一处，列宽等于没设。
        作者没写长短差异时走 `autofit=True`，一列都不动，由内容自己决定。
        """
        share = [max(float(w), 1.0) for w in widths]
        total_share = sum(share)
        total = self._content_width_twips()
        twips = [int(round(total * x / total_share)) for x in share]

        grid = table._tbl.tblGrid
        for col, width in zip(grid.findall(qn("w:gridCol")), twips):
            col.set(qn("w:w"), str(width))

        table.autofit = False  # → w:tblLayout type=fixed
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = parse_xml(f'<w:tblW {nsdecls("w")}/>')
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:type"), "dxa")
        tbl_w.set(qn("w:w"), str(total))

        for row in table.rows:
            for cell, width in zip(row.cells, twips):
                cell.width = Emu(width * 635)  # 1 twip = 635 EMU

    def _mark_header_row(self, table) -> None:
        """表头行跨页重复：`w:tblHeader`。python-docx 没有这项 API，直接写 XML。"""
        from docx.oxml.ns import qn

        tr = table.rows[0]._tr
        trPr = tr.get_or_add_trPr()
        if trPr.find(qn("w:tblHeader")) is None:
            el = trPr.makeelement(qn("w:tblHeader"), {})
            trPr.append(el)

    def _keep_caption_with_figure(self, blocks: list[Block], index: int) -> None:
        """题注与图必须同页，不管题注写在上方还是下方。

        Word 只有「与下段同页」（keepNext）一个方向：题注在上，靠样式里的这一条就够
        （要求文件的「表格题注」就是这么写的）；题注在下，得给**图片段**补。所以按相邻
        关系判断，谁是这一对里的前一件，就给谁加——两个方向都成立，不用作者操心。
        """
        here = self._last_paragraph
        if here is None or index + 1 >= len(blocks):
            return

        def is_figure(node: Block) -> bool:
            return isinstance(node, Figure)

        def is_caption(node: Block) -> bool:
            return isinstance(node, Paragraph) and node.style == "Image Caption"

        this, nxt = blocks[index], blocks[index + 1]
        if not ((is_figure(this) and is_caption(nxt)) or (is_caption(this) and is_figure(nxt))):
            return
        ppr = here._p.get_or_add_pPr()
        if ppr.find(qn("w:keepNext")) is None:
            ppr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))

    def _figure(self, block: Figure) -> None:
        paragraph = self._new_paragraph("Captioned Figure")
        self._last_paragraph = paragraph
        # 稿里的路径相对稿件目录，不按进程 cwd 猜
        path = Path(block.path)
        if not path.is_absolute():
            path = self.base_dir / path
        width = self._clamp(block.width_cm)
        if not is_pdf(path):
            self._embed(paragraph.add_run(), path, width, block.alt)
            return
        # Word 只放得下位图，PDF 里的矢量得先渲染成像素。渲染的临时文件
        # 用完即删：python-docx 在这一步已经把字节读进包里了。
        try:
            png = rasterize(path, width)
        except PdfFigureError as exc:
            raise UnsupportedFeature(t("writer.figure_pdf_failed", self.lang, path=path.name, err=exc)) from exc
        try:
            self._embed(paragraph.add_run(), png, width, block.alt)
        finally:
            png.unlink(missing_ok=True)

    def _clamp(self, width_cm: Optional[float]) -> Optional[float]:
        """超出版心的宽度钳到版心：手册承诺过这一条，不钳就是图和正文对不齐、右侧被裁。"""
        section = self.docx.sections[0] if self.docx.sections else None
        if section is None:
            return width_cm
        widths = (section.page_width, section.left_margin, section.right_margin)
        if any(value is None for value in widths):
            return width_cm
        limit = Emu(widths[0] - widths[1] - widths[2]).cm
        if limit <= 0:
            return width_cm
        return limit if width_cm is None or width_cm > limit else width_cm

    @staticmethod
    def _embed(run, path: Path, width_cm: Optional[float], alt: str = "") -> None:
        shape = (
            run.add_picture(str(path), width=Cm(width_cm))
            if width_cm
            else run.add_picture(str(path))
        )
        if alt:
            shape._inline.docPr.set("descr", alt)

    # ── 内联 ─────────────────────────────────────────────────────────
    def inlines(self, paragraph, nodes: Iterable[Inline]) -> None:
        self._verbatim_style = self._style_name("Verbatim Char")
        for node in _trim_citation_gap(nodes):
            self.inline(paragraph, node)

    def _footnote_reference(self, paragraph, label: str) -> None:
        """脚注引用：run 里挂 `w:footnoteReference`，正文内容稍后统一写进 footnotes.xml。"""
        fids = self._footnote_ids
        if label not in fids:
            fids[label] = len(fids) + 1
        run = paragraph.add_run()
        style = self._style_name("Footnote Reference")
        if style:
            run.style = style
        run._r.append(parse_xml(f'<w:footnoteReference {nsdecls("w")} w:id="{fids[label]}"/>'))

    def _write_footnotes(self) -> None:
        """生成 word/footnotes.xml 并挂上关系。

        python-docx 没有脚注 API，这里直接构造部件。分隔符（id=-1）与续页分隔符（id=0）
        是 Word 的约定，缺了它们脚注区在那条横线上会显示异常。
        """
        body: list[str] = []
        self._footnote_links: list[str] = []
        by_label = {f.label: f for f in self.model.footnotes}
        for label, fid in self._footnote_ids.items():
            footnote = by_label.get(label)
            paragraphs = "".join(self._footnote_paragraph(b) for b in (footnote.blocks if footnote else []))
            if not paragraphs:
                paragraphs = self._footnote_paragraph(Paragraph(inlines=[]))
            body.append(f'<w:footnote w:id="{fid}">{paragraphs}</w:footnote>')

        xml = (
            f'<w:footnotes {nsdecls("w")}>'
            '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
            '<w:footnote w:type="continuationSeparator" w:id="0">'
            "<w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>"
            + "".join(body)
            + "</w:footnotes>"
        )
        part = Part(
            PackURI("/word/footnotes.xml"),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
            b"",
            self.docx.part.package,
        )
        for index, url in enumerate(self._footnote_links):
            rid = part.relate_to(url, RT.HYPERLINK, is_external=True)
            xml = xml.replace(f"{{{{L{index}}}}}", rid)
        part._blob = xml.encode("utf-8")
        self.docx.part.relate_to(part, RT.FOOTNOTES)

    def _footnote_paragraph(self, block) -> str:
        if not isinstance(block, Paragraph):
            raise UnsupportedFeature(t("writer.footnote_paragraph_only", self.lang))
        style = (block.style or "Footnote Text").replace(" ", "")
        runs = "".join(self._run_xml(node) for node in block.inlines)
        # 脚注区每条前面的编号由 `w:footnoteRef` 域生成，缺了它正文有号、注文无号
        number = (
            '<w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr>'
            "<w:footnoteRef/></w:r><w:r><w:t xml:space=\"preserve\"> </w:t></w:r>"
        )
        return f'<w:p><w:pPr><w:pStyle w:val="{style}"/></w:pPr>{number}{runs}</w:p>'

    def _run_xml(self, node) -> str:
        if isinstance(node, LineBreak):
            return "<w:r><w:br/></w:r>"
        if isinstance(node, Code):   # 等宽是叶子节点，别落进下面的纯文本分支
            style = self._style_name("Verbatim Char")
            props = f'<w:rStyle w:val="{style}"/>' if style else ""
            return f'<w:r><w:rPr>{props}</w:rPr><w:t xml:space="preserve">{_escape(node.value)}</w:t></w:r>'
        text = getattr(node, "value", None)
        if text is not None:
            return f'<w:r><w:t xml:space="preserve">{_escape(text)}</w:t></w:r>'
        children = getattr(node, "children", None)
        if children is None:
            raise UnsupportedFeature(t("writer.footnote_node", self.lang,
                                       type=type(node).__name__))
        inner = "".join(self._run_xml(c) for c in children)
        props = ""
        if isinstance(node, Strong):
            props = "<w:b/>"
        elif isinstance(node, Emphasis):
            props = "<w:i/>"
        elif isinstance(node, Strikethrough):
            props = "<w:strike/>"
        elif isinstance(node, Underline):
            props = '<w:u w:val="single"/>'
        elif isinstance(node, Superscript):
            props = '<w:vertAlign w:val="superscript"/>'
        elif isinstance(node, Subscript):
            props = '<w:vertAlign w:val="subscript"/>'
        elif isinstance(node, Link):
            # 脚注区是手工拼的 XML，超链接要先拿到 rId——占位符在这里埋，写部件时替换
            index = len(self._footnote_links)
            self._footnote_links.append(node.url)
            inner = "".join(self._run_xml(child) for child in node.children)
            return f'<w:hyperlink r:id="{{{{L{index}}}}}">{inner}</w:hyperlink>' 
        return inner.replace("<w:r>", f"<w:r><w:rPr>{props}</w:rPr>", 1) if props else inner

    def _render_eq_number(self, chapter: int, index: int) -> str:
        """按模板渲染编号。公式旁的编号与正文里的 `@eq-x` 走的是同一个函数——
        两处必须长得一样，否则就是同一个事实的两个真相源。"""
        fmt = self.eq_number_format
        if "{c}" in fmt:
            if not self.model.meta.chapters:
                raise UnsupportedFeature(t("writer.equation_chapter_without_yaml", self.lang))
            fmt = fmt.replace("{c}", str(chapter))
        return fmt.replace("{n}", str(index))

    def _resolve_crossref(self, target: str) -> str:
        """`@eq-a` → 「式（5-1）」/「Equation (5.1)」：前缀看 YAML 的 crossref，编号按同一份格式。"""
        if target.startswith("eq-"):
            prefix = self._crossref_prefixes.get("eq-prefix", "式")
        elif target.startswith("fig-"):
            prefix = self._crossref_prefixes.get("fig-prefix", "图")
        elif target.startswith("tbl-"):
            prefix = self._crossref_prefixes.get("tbl-prefix", "表")
        else:
            raise UnsupportedFeature(t("writer.crossref_type", self.lang, target=target))
        parts = self._eq_numbers.get(target)
        if parts is None:
            raise UnsupportedFeature(t("writer.crossref_missing", self.lang, target=target))
        return f"{prefix}{self._render_eq_number(*parts)}"

    def _new_paragraph(self, style: str | None = None):
        """按样式名新建段落；赋样式对象可省去 python-docx 的按名查表。"""
        paragraph = self.docx.add_paragraph()
        if style:
            obj = self._style_object(style) or self._style_object("Normal")
            if obj is not None:
                paragraph.style = obj
        return paragraph

    def _style_object(self, name: str):
        """样式名 → 样式对象；比对不分大小写（Word 内置样式名存为小写 `heading 1`）。"""
        if self._style_objects is None:
            self._style_objects = {style.name.casefold(): style for style in self.docx.styles}
        return self._style_objects.get(name.casefold())

    def _style_name(self, name: str) -> str | None:
        """按名称取样式名（没有就返回 None，不静默套用别的样式）。"""
        if self._style_names is None:
            self._style_names = {style.name.casefold(): style.name for style in self.docx.styles}
        return self._style_names.get(name.casefold())

    def _style_ref(self, name: str) -> str:
        """样式名或 styleId → 该样式在本 docx 里的**名字**。

        `add_paragraph(style=…)` 按名字查；传 styleId（如 `SourceCode`）时 python-docx
        会退回按 id 查并发一条弃用警告。这里先换成名字，调用处不刷警告。
        """
        if self._style_refs is None:
            self._style_refs = {}
            for style in self.docx.styles:
                self._style_refs.setdefault(style.name.casefold(), style.name)
                self._style_refs.setdefault(style.style_id.casefold(), style.name)
        return self._style_refs.get(name.casefold(), name)

    def inline(self, paragraph, node: Inline) -> None:
        if isinstance(node, Code):
            run = paragraph.add_run(node.value)
            if self._verbatim_style:
                run.style = self._verbatim_style
        elif isinstance(node, Strong):
            self._styled(paragraph, node.children, bold=True)
        elif isinstance(node, Emphasis):
            self._styled(paragraph, node.children, italic=True)
        elif isinstance(node, Strikethrough):
            self._styled(paragraph, node.children, strike=True)
        elif isinstance(node, Underline):
            self._styled(paragraph, node.children, underline=True)
        elif isinstance(node, Superscript):
            self._styled(paragraph, node.children, superscript=True)
        elif isinstance(node, Subscript):
            self._styled(paragraph, node.children, subscript=True)
        elif isinstance(node, LineBreak):
            paragraph.add_run().add_break()
        elif isinstance(node, Math):
            try:
                nodes = latex_to_omml(node.latex, self.lang)
            except UnsupportedMath as exc:
                snippet = node.latex.strip()[:30]
                raise UnsupportedFeature(t("writer.inline_equation_failed", self.lang,
                                           snippet=snippet, err=exc)) from exc
            paragraph._p.append(self._omath(nodes))
        elif isinstance(node, Link):
            self._hyperlink(paragraph, node)
        elif isinstance(node, CrossRef):
            self._trim_cjk_gap(paragraph)
            paragraph.add_run(self._resolve_crossref(node.target))
        elif isinstance(node, Citation):
            if self.citations is None:
                raise UnsupportedFeature(t("writer.citation_without_bibliography", self.lang))
            keys = [(item.key, item.locator, item.prefix, item.label) for item in node.items]
            segments = self.citations.cite(keys, narrative=node.narrative)
            # 中文行文里「基石 [1]」应为「基石[1]」；西文保留空格
            if not node.narrative:
                self._trim_cjk_gap(paragraph)
            if self.model.meta.link_citations:
                self._linked_run(paragraph, segments, f"_Ref_{keys[0][0]}")
            else:
                for text, fmt in segments:
                    self._run_element(paragraph._p, text, fmt)
        elif isinstance(node, FootnoteRef):
            self._footnote_reference(paragraph, node.label)
        else:
            paragraph.add_run(getattr(node, "value", str(node)))

    @staticmethod
    def _trim_cjk_gap(paragraph) -> None:
        """中文后面的那个空格要去掉：引文、交叉引用都走同一条规则。"""
        if not paragraph.runs:
            return
        last = paragraph.runs[-1]
        if last.text.endswith(" ") and _ends_with_cjk(last.text):
            last.text = last.text.rstrip() or " "

    def _linked_run(self, paragraph, segments, anchor: str) -> None:
        """引文跳到文末对应条目（Word 内链）。片段里的格式（Nature 的上标编号）要保留。"""
        element = parse_xml(f'<w:hyperlink {nsdecls("w")} w:anchor="{anchor}"/>')
        for text, fmt in segments:
            element.append(self._run_element(element, text, fmt))
        paragraph._p.append(element)

    @staticmethod
    def _run_element(parent, text: str, fmt: dict):
        """在给定父元素下造一个带格式的 run；hyperlink / 普通段落共用。"""
        run = parent.makeelement(qn("w:r"), {})
        if fmt:
            rpr = run.makeelement(qn("w:rPr"), {})
            for name, tag in (("bold", "b"), ("italic", "i"), ("strike", "strike")):
                if fmt.get(name):
                    rpr.append(run.makeelement(qn(f"w:{tag}"), {}))
            if fmt.get("superscript"):
                el = run.makeelement(qn("w:vertAlign"), {})
                el.set(qn("w:val"), "superscript")
                rpr.append(el)
            if fmt.get("subscript"):
                el = run.makeelement(qn("w:vertAlign"), {})
                el.set(qn("w:val"), "subscript")
                rpr.append(el)
            run.append(rpr)
        t = run.makeelement(qn("w:t"), {})
        t.text = text
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        run.append(t)
        return run

    def _hyperlink(self, paragraph, node: Link) -> None:
        """超链接是关系（r:id）+ `w:hyperlink`，不是普通 run。"""
        from docx.oxml.ns import nsdecls as _nsdecls

        r_id = paragraph.part.relate_to(node.url, RT.HYPERLINK, is_external=True)
        element = parse_xml(
            f'<w:hyperlink {_nsdecls("w", "r")} r:id="{r_id}"/>'
        )
        style = self._style_name("Hyperlink")
        for text, fmt in self._flatten_text(node.children):
            run = element.makeelement(qn("w:r"), {})
            rpr = run.makeelement(qn("w:rPr"), {})
            if style:
                rpr.append(run.makeelement(qn("w:rStyle"), {qn("w:val"): style.replace(" ", "")}))
            for name in ("bold", "italic", "strike", "underline"):
                if fmt.get(name):
                    rpr.append(run.makeelement(qn(f"w:{_W_FMT[name]}"), {}))
            if len(rpr):
                run.append(rpr)
            t = run.makeelement(qn("w:t"), {})
            t.text = text
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            run.append(t)
            element.append(run)
        paragraph._p.append(element)

    def _prose_text(self) -> str:
        """正文文字的前一小段：判语种用，不需要全量。"""
        parts: list[str] = []
        for block in self.model.blocks:
            if isinstance(block, (Heading, Paragraph)):
                parts.append(_plain_text(block.inlines))
            if sum(len(p) for p in parts) > 800:
                break
        return "\n".join(parts)

    def _flatten_text(self, nodes, fmt=None):
        """把带格式的行内树摊平成 (文本, 格式) 序列，供超链接这类要自己造 run 的地方用。"""
        fmt = dict(fmt or {})
        for node in nodes:
            if isinstance(node, Code):
                yield node.value, fmt
            elif isinstance(node, LineBreak):
                continue  # 链接文字里的换行没有落点，跳过
            elif type(node) in self._FORMAT_ATTR:
                yield from self._flatten_text(node.children, {**fmt, self._FORMAT_ATTR[type(node)]: True})
            else:
                text = getattr(node, "value", None)
                if text is None:
                    raise UnsupportedFeature(t("writer.link_child", self.lang,
                                               type=type(node).__name__))
                yield text, fmt

    _FORMAT_ATTR = {
        Strong: "bold",
        Emphasis: "italic",
        Strikethrough: "strike",
        Underline: "underline",
        Superscript: "superscript",
        Subscript: "subscript",
    }

    def _styled(self, paragraph, children, **fmt) -> None:
        """把一层格式作用到子节点上；子节点还是容器就递归合并——`***粗斜体***` 要能落成
        一个既粗又斜的 run，而不是报错。"""
        for child in children:
            if isinstance(child, Code):
                run = paragraph.add_run(child.value)
                if self._verbatim_style:
                    run.style = self._verbatim_style
                self._apply_font(run, fmt)
            elif type(child) in self._FORMAT_ATTR:
                self._styled(paragraph, child.children, **{**fmt, self._FORMAT_ATTR[type(child)]: True})
            elif isinstance(child, Math):
                paragraph._p.append(self._omath(latex_to_omml(child.latex, self.lang)))
            elif hasattr(child, "children"):
                raise UnsupportedFeature(t("writer.inline_format_child", self.lang,
                                           type=type(child).__name__))
            else:
                text = getattr(child, "value", None)
                if text is None:
                    raise UnsupportedFeature(t("writer.inline_node", self.lang,
                                               type=type(child).__name__))
                run = paragraph.add_run(text)
                self._apply_font(run, fmt)

    @staticmethod
    def _apply_font(run, fmt: dict) -> None:
        for key, value in fmt.items():
            setattr(run.font, key, value)


def write_docx(
    model: Document,
    template: Path | str,
    output: Path | str,
    citations=None,
    base_dir: Path | str | None = None,
    equation_number_placement: str = "right",
    equation_number_format: str = "",
    no_row_split: bool = False,
    repeat_header: bool = False,
    lang: Optional[str] = None,
) -> Path:
    """把文档模型写成 .docx。模板提供 styles.xml（由格式要求文件生成）。

    `citations` 是 `transform.citations.CitationEngine`：有文献库时传进来，文献引用与
    文末列表都由它渲染；不传时遇到 citation 会明确失败。
    `base_dir` 是稿件所在目录：稿里的图片路径都是相对它的，不能按进程 cwd 解析——
    否则从别的目录调用时图全找不到（预览与流水线走的是同一个入口）。
    `equation_number_placement`：`right` 编号靠版心右边（默认），`inline` 紧跟公式。
    `equation_number_format`：编号外观模板（`（{c}-{n}）`）；空串按语种给默认。
    """
    docx = open_docx(str(template))
    writer = _Writer(
        model,
        docx,
        citations=citations,
        base_dir=base_dir,
        equation_number_placement=equation_number_placement,
        equation_number_format=equation_number_format,
        no_row_split=no_row_split,
        repeat_header=repeat_header,
        lang=lang,
    )
    writer.write()
    out = Path(output)
    docx.save(str(out))
    # 时间戳钉死：同一份稿件重复导出要逐字节一致，否则 Git 里每次都是「内容没改却变了」
    normalize_zip(out)
    return out

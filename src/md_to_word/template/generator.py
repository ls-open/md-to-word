#!/usr/bin/env python3
"""
生成学术 Word 模板 (docx)

从纯净现代 Word 365 基础母版出发，根据 requirements.txt 格式要求修改样式，
输出是只含样式与页面设置的 template.docx，写出器直接用它。

用法:
    python3 generate_template.py requirements.txt -o template.docx
    python3 generate_template.py --help
"""

import os
import re
import shutil
import sys
import tempfile
from dataclasses import dataclass, field
from typing import Optional
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from lxml import etree

from md_to_word.config import get_base_template_path
from md_to_word.i18n import t
from md_to_word.utils.skill_utils import (
    TOC_MARKER_RE,
    clear_color_theme_attrs,
    clear_font_theme_attrs,
    get_logger,
    strip_theme_attrs_from_xml,
)

logger = get_logger(__name__)

# ── XML 命名空间 ──
NS_W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS_M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'

def W(tag):
    return f'{{{NS_W}}}{tag}'

def M(tag):
    return f'{{{NS_M}}}{tag}'

# ── 字号映射（中文字号 → pt）──
SIZE_MAP = {
    '初号': Pt(42),
    '小初': Pt(36),
    '一号': Pt(26),
    '小一': Pt(24),
    '二号': Pt(22),
    '小二': Pt(18),
    '三号': Pt(16),
    '小三': Pt(15),
    '四号': Pt(14),
    '小四': Pt(12),
    '五号': Pt(10.5),
    '小五': Pt(9),
    '六号': Pt(7.5),
    '小六': Pt(6.5),
    '七号': Pt(5.5),
    '八号': Pt(5),
}

# ── 样式名映射（封闭集合：除 Title 兼容「文档标题 / 论文标题」外，其余每种样式只认唯一中文名）──
STYLE_MAP = {
    '论文标题': 'Title',
    '文档标题': 'Title',
    '副标题': 'Subtitle',
    '作者': 'Author',
    '日期': 'Date',
    '摘要': 'Abstract',
    '摘要标题': 'Abstract Title',
    '关键词': 'Keywords',
    '一级标题': 'Heading 1',
    '二级标题': 'Heading 2',
    '三级标题': 'Heading 3',
    '四级标题': 'Heading 4',
    '五级标题': 'Heading 5',
    '六级标题': 'Heading 6',
    '正文': 'Body Text',
    '首段': 'First Paragraph',
    '普通': 'Normal',
    '表格题注': 'Table Caption',
    '图片题注': 'Image Caption',
    '图': 'Captioned Figure',
    '表注': 'Table Note',
    '参考文献': 'Bibliography',
    '参考文献标题': 'References Title',
    '紧凑': 'Compact',
    '引用块': 'Block Text',
    # 代码块（``` 围栏）的段落样式。底模已给一套等宽字体 + 无缩进的默认值，
    # 所以「代码：」是可写可不写的那一类——写它就整行覆盖掉那套默认。
    '代码': 'SourceCode',
    # 脚注：写法与别的样式行一样，但它是**可写可不写**的那一类——稿件里有 `[^1]` 才用得上，
    # 所以不进必填清单。不写时由 apply_footnote_format 从正文推一套默认补上（见那里的说明）。
    '脚注': 'Footnote Text',
    # 页眉/页脚是底模自带的段落样式（w:name = header / footer），
    # 于是「页眉：宋体 小五 居中」就是一条标准样式行，复用全套段落属性。
    # 注意：这里设的是**段落样式**，页眉/页脚的**内容**另由「页眉内容：」等行给。
    '页眉': 'Header',
    '页脚': 'Footer',
    # 目录：标题与条目分属不同样式，两类都会静默建空壳（不写就字体字号全落空），故都列。
    # 条目样式对应 Word 内置的 TOC 1/2/3，由目录域自己按级套用——模板里没有这三条，
    # 生成的目录就会用 Word 内置默认（字体、缩进、间距都不是要求的那套）。
    '目录标题': 'TOC Heading',
    '目录一级': 'TOC 1',
    '目录二级': 'TOC 2',
    '目录三级': 'TOC 3',
}

STYLE_CN = {v: k for k, v in STYLE_MAP.items()}

# ── 页眉/页脚内容行（键 → (页眉|页脚, default|even|first)）──
# default = 奇偶未区分时的那一套；even = 偶数页；first = 首页（首页不同）
DECOR_CONTENT_KEYS = {
    '页眉内容': ('header', 'default'),
    '偶数页页眉内容': ('header', 'even'),
    '首页页眉内容': ('header', 'first'),
    '页脚内容': ('footer', 'default'),
    '偶数页页脚内容': ('footer', 'even'),
    '首页页脚内容': ('footer', 'first'),
}

# ── 页码格式（唯一写法）──
PAGE_NUMBER_FMT = {
    '阿拉伯': 'decimal',
    '罗马': 'upperRoman',
    '小写罗马': 'lowerRoman',
    '字母': 'lowerLetter',
    '大写字母': 'upperLetter',
}


# ── 纸张尺寸（唯一写法 → 内部键）──
PAGE_SIZE_MAP = {'A4': 'a4', 'Letter': 'letter'}

# 纸张的**唯一一张尺寸表**（cm）。版心宽度、交给 Word 的 pgSz、导出后对账的期望值都由它算。
PAGE_SIZES_CM = {'a4': (21.0, 29.7), 'letter': (21.59, 27.94)}
_TWIPS_PER_CM = 566.93


def page_size_twips(page_size: str) -> tuple:
    """纸张 → (宽, 高) twips（对账侧与写入侧共用，避免各写一份常量）。"""
    w_cm, h_cm = PAGE_SIZES_CM.get(page_size, PAGE_SIZES_CM['a4'])
    return int(round(w_cm * _TWIPS_PER_CM)), int(round(h_cm * _TWIPS_PER_CM))

# ── 页眉/页脚内容里可用的域标记（唯一写法；未知标记报错，不静默漏掉）──
#   {页码}           → PAGE 域（当前页）
#   {总页数}         → NUMPAGES 域
#   {样式标题:S}     → STYLEREF 域，取当前页上第一个 S 样式段落的文字
#                      （如 {样式标题:Heading 1} = 当前章标题）
_RE_DECOR_TOKEN = re.compile(r'\{([^{}]*)\}')
_RE_STYLEREF = re.compile(r'^样式标题\s*[:：]\s*(.+?)\s*$')

# ── 对齐方式（唯一写法：居中/左对齐/右对齐/两端对齐）──
# **唯一一张表**：中文写法 → (python-docx 枚举, OOXML 值)。下面两个视图与
# audit_export 里那份都从这里派生，不再各写一遍。
_ALIGN = {
    '居中': (WD_ALIGN_PARAGRAPH.CENTER, 'center'),
    '左对齐': (WD_ALIGN_PARAGRAPH.LEFT, 'left'),
    '右对齐': (WD_ALIGN_PARAGRAPH.RIGHT, 'right'),
    '两端对齐': (WD_ALIGN_PARAGRAPH.JUSTIFY, 'both'),
}
ALIGN_MAP = {cn: enum for cn, (enum, _) in _ALIGN.items()}          # 解析用
ALIGN_VAL = {cn: val for cn, (_, val) in _ALIGN.items()}            # 写 XML / 导出后对账

# ── 已知 CJK / Latin 字体（用于分离中西文）──
CJK_FONTS = {'宋体', '黑体', '仿宋', 'FangSong', '楷体', '微软雅黑'}
LATIN_FONTS = {'Times New Roman', 'Arial', 'Calibri', 'Verdana'}


def ensure_rpr(element):
    """确保元素有 rPr 子元素，返回 rPr"""
    rpr = element.find(W('rPr'))
    if rpr is None:
        rpr = etree.SubElement(element, W('rPr'))
    return rpr


# ── OOXML 子元素顺序 ──
# Word 对 w:rPr / w:pPr / w:tblStylePr 的子元素顺序有 schema 强制要求，
# 顺序错了 Word 会**整段忽略**（不报错）。按位置插入，不无脑 append。
_RPR_ORDER = ['rStyle', 'rFonts', 'b', 'bCs', 'i', 'iCs', 'caps', 'smallCaps', 'strike',
              'dstrike', 'outline', 'shadow', 'emboss', 'imprint', 'noProof', 'snapToGrid',
              'vanish', 'webHidden', 'color', 'spacing', 'w', 'kern', 'position', 'sz', 'szCs',
              'highlight', 'u', 'effect', 'bdr', 'shd', 'fitText', 'vertAlign', 'rtl', 'cs',
              'em', 'lang', 'eastAsianLayout', 'specVanish', 'oMath']
_PPR_ORDER = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl',
              'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens',
              'kinsoku', 'wordWrap', 'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN',
              'bidi', 'adjustRightInd', 'snapToGrid', 'spacing', 'ind', 'contextualSpacing',
              'mirrorIndents', 'suppressOverlap', 'jc', 'textDirection', 'textAlignment',
              'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr', 'sectPr', 'pPrChange']
_STYLE_ORDER = ['name', 'aliases', 'basedOn', 'next', 'link', 'autoRedefine', 'hidden',
                'uiPriority', 'semiHidden', 'unhideWhenUsed', 'qFormat', 'locked', 'personal',
                'personalCompose', 'personalReply', 'rsid', 'pPr', 'rPr', 'tblPr', 'trPr',
                'tcPr', 'tblStylePr']
_TBLSTYLEPR_ORDER = ['pPr', 'rPr', 'tblPr', 'trPr', 'tcPr']
_TBLPR_ORDER = ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize',
                'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders',
                'shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption', 'tblDescription']
_TCPR_ORDER = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap',
               'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark']
_TRPR_ORDER = ['cnfStyle', 'divId', 'gridBefore', 'gridAfter', 'wBefore', 'wAfter', 'cantSplit',
               'trHeight', 'tblHeader', 'tblCellSpacing', 'jc', 'hidden']
_SECTPR_ORDER = ['footnotePr', 'endnotePr', 'type', 'pgSz', 'pgMar', 'paperSrc', 'pgBorders',
                 'lnNumType', 'pgNumType', 'cols', 'formProt', 'vAlign', 'noEndnote',
                 'titlePg', 'textDirection', 'bidi', 'rtlGutter', 'docGrid', 'printerSettings']
# ── 段落下边框（页眉下的那条线就靠它）──
# 「双线」用 OOXML 的 thinThickMediumGap——Word 画成一粗一细两条，
# 学位论文的页眉线即此，常见取值 1.5 磅（`sz=12`）。
BOTTOM_BORDER_STYLE = {'single': 'single', 'double': 'thinThickMediumGap'}
BOTTOM_BORDER_DEFAULT_PT = {'single': 0.5, 'double': 1.5}

# ── 版心宽度（twips）──
# 目录的「点线制表位」要把页码顶到版心右边界，位置得由纸张与页边距算，不能写死：
# 换纸张或改页边距后，写死的制表位会顶出边界或提前断线。
# 兜底值 = A4 − 左右各 2.54cm；正常路径由 generate_template 按实际几何注入
_FALLBACK_CONTENT_WIDTH_TWIPS = int(round((PAGE_SIZES_CM['a4'][0] - 2.54 * 2) * _TWIPS_PER_CM))


def _content_width_twips(page_size, margins, gutter_cm=None) -> int:
    """正文可用宽度（twips）。装订线是折进**左**边距写的，所以要一起减掉。"""
    width_cm = PAGE_SIZES_CM.get(page_size, PAGE_SIZES_CM['a4'])[0]
    left = margins['left'] + (gutter_cm or 0)
    return int(round((width_cm - left - margins['right']) * _TWIPS_PER_CM))


def body_line_spacing_multiple(styles) -> float:
    """正文的行距**倍数**——文档网格补偿按它算。

    `styles` 是任一含 `style_name` / `line_spacing` 的序列（写入侧是 `spec.styles`，
    读取侧是解析出来的样式行）。优先「正文」，退回「普通」；固定值行距（`行距20磅`）
    不是倍数、不参与相乘（Word 只把行高向上取整到网格），故返回 1.0。
    """
    for st in styles:
        if st.get('style_name') in ('Body Text', 'Normal'):
            ls = st.get('line_spacing')
            if isinstance(ls, float) and ls > 0:
                return ls
    return 1.0


def _body_height_twips(page_size, margins) -> int:
    _, h_cm = PAGE_SIZES_CM.get(page_size, PAGE_SIZES_CM['a4'])
    return int(round((h_cm - margins['top'] - margins['bottom']) * _TWIPS_PER_CM))


def doc_grid_pitch(spec, margins) -> int:
    """`文档网格：每页N行` → `w:docGrid/@w:linePitch`（每行多少 twips）。

    写入侧、导出后对账侧共用这一份，免得两边各算一遍、算法漂移还看不出来。
    """
    # **行距的「倍数」会把网格间距再乘一遍**：声明 每页22行 时，配 1.5 倍行距实测只落
    # 14 行、配 2.0 倍只落 11 行（正好是 22÷倍数）。作者要的是「每页 N 行」，就把这个
    # 倍数除掉，让**结果**真的是 N 行——而不是反过来要求他改行距。
    return int(_body_height_twips(spec.page_size, margins)
               / spec.doc_grid_lines / body_line_spacing_multiple(spec.styles))


def doc_grid_lines_from_pitch(pitch: int, line_multiple: float, page_size, margins) -> int:
    """`w:docGrid/@linePitch` → 声明的「每页 N 行」。是 `doc_grid_pitch` 的逆，
    读参考 .docx 时用（把别人模板里的网格反推成本技能的写法）。"""
    return max(1, int(round(_body_height_twips(page_size, margins) / pitch / line_multiple)))

# python-docx 枚举 → OOXML 值（由上面那张中文表派生）
_ALIGN_VAL = {enum: val for enum, val in _ALIGN.values()}


def ensure_ordered_child(parent, tag, order):
    """按 schema 顺序取/建子元素。顺序表外的标签追加在末尾。"""
    el = parent.find(W(tag))
    if el is not None:
        return el
    el = etree.Element(W(tag))
    idx = order.index(tag) if tag in order else len(order)
    for i, ch in enumerate(parent):
        ln = etree.QName(ch).localname
        if ln in order and order.index(ln) > idx:
            parent.insert(i, el)
            return el
    parent.append(el)
    return el


def split_fonts(font_names, default_latin=None):
    """把字体名列表拆成 (中日韩字体, 西文字体)"""
    cjk = None
    latin = None
    for f in (font_names or []):
        if f in CJK_FONTS or re.search(r'[一-鿿]', f):
            cjk = f
        else:
            latin = f
    if not latin and default_latin:
        latin = default_latin
    return cjk, latin


def _neutralize_inherited_first_line(ind):
    """设置悬挂缩进时，显式将继承的首行缩进置零，避免二者冲突。"""
    ind.set(W('firstLineChars'), '0')
    ind.attrib.pop(W('firstLine'), None)


def write_run_props(rpr, props, cjk=None, latin=None):
    """写入字符级属性（字体、字号、加粗、倾斜、颜色）。"""
    if cjk or latin:
        rfonts = ensure_ordered_child(rpr, 'rFonts', _RPR_ORDER)
        if cjk:
            rfonts.set(W('eastAsia'), cjk)
        if latin:
            rfonts.set(W('ascii'), latin)
            rfonts.set(W('hAnsi'), latin)
        clear_font_theme_attrs(rfonts)

    if props.get('font_size'):
        half_pt = str(int(round(props['font_size'].pt * 2)))
        ensure_ordered_child(rpr, 'sz', _RPR_ORDER).set(W('val'), half_pt)
        # 同步复杂文字字号（szCs），保持多语言环境下字号一致
        sz_cs = rpr.find(W('szCs'))
        if sz_cs is not None:
            sz_cs.set(W('val'), half_pt)

    # 显式设置加粗状态（开/关），避免底模样式的加粗属性隐式继承
    if props.get('bold') is not None:
        b = ensure_ordered_child(rpr, 'b', _RPR_ORDER)
        if not props['bold']:
            b.set(W('val'), '0')
        else:
            b.attrib.pop(W('val'), None)
        b_cs = rpr.find(W('bCs'))
        if b_cs is not None:
            if not props['bold']:
                b_cs.set(W('val'), '0')
            else:
                b_cs.attrib.pop(W('val'), None)

    # 显式设置倾斜状态（开/关），避免底模样式的倾斜属性隐式继承
    if props.get('italic') is not None:
        it = ensure_ordered_child(rpr, 'i', _RPR_ORDER)
        if not props['italic']:
            it.set(W('val'), '0')
        else:
            it.attrib.pop(W('val'), None)
        it_cs = rpr.find(W('iCs'))
        if it_cs is not None:
            if not props['italic']:
                it_cs.set(W('val'), '0')
            else:
                it_cs.attrib.pop(W('val'), None)

    # 统一设置黑色并清除主题色，同时清除底模可能残留的字符间距加宽（spacing）与字距调整（kern）
    color = ensure_ordered_child(rpr, 'color', _RPR_ORDER)
    color.set(W('val'), '000000')
    clear_color_theme_attrs(color)
    for tag in ('spacing', 'kern'):
        el = rpr.find(W(tag))
        if el is not None:
            rpr.remove(el)
    # 作者写了「字距N磅」才写回去（val 单位是 1/20 磅）
    if props.get('char_spacing_pt'):
        sp = ensure_ordered_child(rpr, 'spacing', _RPR_ORDER)
        sp.set(W('val'), str(int(round(props['char_spacing_pt'] * 20))))


def write_para_props(ppr, props):
    """把**段落级**属性写进任意 pPr 容器（样式 pPr，或 w:tblStylePr/pPr）。"""
    if props.get('alignment') is not None:
        jc = ensure_ordered_child(ppr, 'jc', _PPR_ORDER)
        jc.set(W('val'), _ALIGN_VAL.get(props['alignment'], 'left'))

    spacing = None
    if props.get('line_spacing') is not None:
        spacing = ensure_ordered_child(ppr, 'spacing', _PPR_ORDER)
        ls = props['line_spacing']
        if isinstance(ls, MinLineSpacing):
            spacing.set(W('line'), str(int(round(ls.pt * 20))))
            spacing.set(W('lineRule'), 'atLeast')
        elif isinstance(ls, Pt):
            spacing.set(W('line'), str(int(round(ls.pt * 20))))
            spacing.set(W('lineRule'), 'exact')
        else:
            spacing.set(W('line'), str(int(round(float(ls) * 240))))
            spacing.set(W('lineRule'), 'auto')

    if 'space_before' in props:
        spacing = spacing if spacing is not None else ensure_ordered_child(ppr, 'spacing', _PPR_ORDER)
        spacing.set(W('before'), str(int(round(props['space_before'].pt * 20))))
    if 'space_after' in props:
        spacing = spacing if spacing is not None else ensure_ordered_child(ppr, 'spacing', _PPR_ORDER)
        spacing.set(W('after'), str(int(round(props['space_after'].pt * 20))))

    if props.get('left_indent_chars') is not None:
        ind = ensure_ordered_child(ppr, 'ind', _PPR_ORDER)
        ind.set(W('leftChars'), str(int(props['left_indent_chars'] * 100)))
        ind.attrib.pop(W('left'), None)

    if props.get('dot_leader_tab'):
        tabs = ensure_ordered_child(ppr, 'tabs', _PPR_ORDER)
        for old in tabs.findall(W('tab')):
            tabs.remove(old)
        tab = etree.SubElement(tabs, W('tab'))
        tab.set(W('val'), 'right')
        tab.set(W('leader'), 'dot')
        tab.set(W('pos'), str(props.get('dot_leader_pos', _FALLBACK_CONTENT_WIDTH_TWIPS)))

    if props.get('join_toc'):
        outline = ensure_ordered_child(ppr, 'outlineLvl', _PPR_ORDER)
        outline.set(W('val'), '0')

    if props.get('keep_next'):
        ensure_ordered_child(ppr, 'keepNext', _PPR_ORDER)
    if props.get('keep_lines'):
        ensure_ordered_child(ppr, 'keepLines', _PPR_ORDER)
    if props.get('page_break_before'):
        ensure_ordered_child(ppr, 'pageBreakBefore', _PPR_ORDER)
    if props.get('widow_control') is not None:
        wc = ensure_ordered_child(ppr, 'widowControl', _PPR_ORDER)
        wc.set(W('val'), '1' if props['widow_control'] else '0')

    if props.get('no_border') or props.get('bottom_border') == 'none':
        p_bdr = ppr.find(W('pBdr'))
        if p_bdr is not None:
            ppr.remove(p_bdr)

    if props.get('bottom_border') in ('single', 'double'):
        p_bdr = ppr.find(W('pBdr'))
        if p_bdr is None:
            p_bdr = ensure_ordered_child(ppr, 'pBdr', _PPR_ORDER)
        old = p_bdr.find(W('bottom'))
        if old is not None:
            p_bdr.remove(old)
        bot = etree.SubElement(p_bdr, W('bottom'))
        bot.set(W('val'), BOTTOM_BORDER_STYLE[props['bottom_border']])
        bot.set(W('sz'), _pt_to_border_sz(props.get('border_width_pt', BOTTOM_BORDER_DEFAULT_PT[props['bottom_border']])))
        bot.set(W('space'), '1')
        bot.set(W('color'), 'auto')

    if props.get('indent') is not None:
        ind = ensure_ordered_child(ppr, 'ind', _PPR_ORDER)
        indent_type = props['indent_type']
        indent_val = props['indent']
        if props.get('indent_unit', 'pt') == 'chars':
            # 字符单位用 firstLineChars/hangingChars，Word 根据段落字号自动换算实际缩进，
            # 字号变化时比例自动保持；与 firstLine（绝对 pt）不同，故两者互斥地清掉一个。
            if indent_type == 'first_line':
                ind.set(W('firstLineChars'), str(int(indent_val * 100)))
                ind.attrib.pop(W('firstLine'), None)
            elif indent_type == 'hanging':
                ind.set(W('hangingChars'), str(int(indent_val * 100)))
                ind.set(W('leftChars'), str(int(indent_val * 100)))
                _neutralize_inherited_first_line(ind)
                ind.attrib.pop(W('hanging'), None)
                ind.attrib.pop(W('left'), None)
        else:
            twips = str(int(round(indent_val * 20)))  # indent_val 单位为 pt
            if indent_type == 'first_line':
                ind.set(W('firstLine'), twips)
                ind.attrib.pop(W('firstLineChars'), None)
            elif indent_type == 'hanging':
                ind.set(W('hanging'), twips)
                ind.set(W('left'), twips)
                _neutralize_inherited_first_line(ind)
                ind.attrib.pop(W('hangingChars'), None)
                ind.attrib.pop(W('leftChars'), None)


# ── 解析函数 ──

def parse_font_name(text):
    """提取字体名，返回 [cjk_font, latin_font] 或 None"""
    parts = re.split(r'[/／]', text)
    cjk = None
    latin = None

    for part in parts:
        p = part.strip()
        if not p:
            continue

        known_match = None
        for known in sorted(CJK_FONTS | LATIN_FONTS, key=len, reverse=True):
            if re.search(r'\b' + re.escape(known) + r'\b', p):
                known_match = known
                break

        if known_match:
            if known_match in CJK_FONTS:
                cjk = known_match
            else:
                latin = known_match
        else:
            candidate = re.split(r'[,，]', p)[0].strip()
            # 字体名里得有字（字母或汉字）：纯数字（如把字号写成 `12` 漏了 pt）不是字体名，
            # 放它过去就等于把"写错的字号"吞成字体名。
            if candidate and not re.match(r'\d', candidate) \
                    and re.search(r'[^\W\d_]|[\u4e00-\u9fff]', candidate):
                if re.search(r'[\u4e00-\u9fff]', candidate):
                    cjk = candidate
                else:
                    latin = candidate

    result = []
    if cjk:
        result.append(cjk)
    if latin:
        result.append(latin)
    return result if result else None


class MinLineSpacing:
    """行距「最小值 N 磅」（`w:lineRule="atLeast"`）：至少 N 磅，内容更高时自动撑开。

    Pt 表示固定值、float 表示倍数，三种不能用同一个类型表达——固定值会被超高的内容裁掉，
    与最小值在单行时长得一样，只有折行才分得出，故类型上就必须分开。
    """
    __slots__ = ('pt',)

    def __init__(self, pt: float):
        self.pt = float(pt)

    def __eq__(self, other):
        return isinstance(other, MinLineSpacing) and other.pt == self.pt

    def __repr__(self):
        return f'MinLineSpacing({self.pt}磅)'


class TemplateError(ValueError):
    """格式要求或模板无法生成。

    作为库调用时（`run_export_pipeline`、`build_template`）必须抛这个异常，不能
    `sys.exit`：调用方的进程不该被一份写错的 requirements 直接终结。
    """


# ── 样式操作函数 ──

# Word 内置的目录条目样式：styleId = TOC1/TOC2/TOC3，名字是小写的 "toc 1/2/3"，
# 且**不能带 `w:customStyle="1"`**。带了这个标记，Word 就把它们当用户自建样式，
# 更新目录域时会按自己的内置默认重建，我们写的字体/字号/加粗全部丢失。
_BUILTIN_TOC = {"TOC 1": ("TOC1", "toc 1"), "TOC 2": ("TOC2", "toc 2"), "TOC 3": ("TOC3", "toc 3")}


def _ensure_builtin_toc_style(style, style_name: str) -> None:
    """把目录条目样式写成 Word 的**内置**样式，而不是自建样式。"""
    builtin = _BUILTIN_TOC.get(style_name)
    if builtin is None:
        return
    style_id, name = builtin
    element = style.element
    element.attrib.pop(W('customStyle'), None)
    element.set(W('styleId'), style_id)
    name_el = element.find(W('name'))
    if name_el is None:
        name_el = etree.SubElement(element, W('name'))
    name_el.set(W('val'), name)


def modify_style(doc, style_name, props, default_latin=None):
    """修改或创建样式的属性"""
    try:
        style = doc.styles[style_name]
    except KeyError:
        style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        try:
            normal_style = doc.styles['Normal']
            style.base_style = normal_style
        except KeyError:
            pass
        logger.info(f' ✓ {style_name}: 已创建新样式（basedOn Normal）')

    _ensure_builtin_toc_style(style, style_name)

    cjk, latin = split_fonts(props.get('font_name'), default_latin)
    if not cjk and not latin and default_latin:
        latin = default_latin

    rpr = ensure_rpr(style.element)
    ppr = ensure_ordered_child(style.element, 'pPr', _STYLE_ORDER)

    if props.get('no_border') and ppr.find(W('pBdr')) is not None:
        logger.info(f' ✓ {style_name}: 已去除段落边框')

    write_run_props(rpr, props, cjk=cjk, latin=latin)
    write_para_props(ppr, props)

    # 同步更新关联的字符样式（w:link），保证段落与字符级属性一致
    link_el = style.element.find(W('link'))
    if link_el is not None:
        link_id = link_el.attrib.get(W('val'))
        parent = style.element.getparent()
        if link_id and parent is not None:
            for s in parent.findall(W('style')):
                if s.attrib.get(W('styleId')) == link_id:
                    link_rpr = ensure_rpr(s)
                    write_run_props(link_rpr, props, cjk=cjk, latin=latin)
                    break

    logger.info(f' ✓ {style_name}: 已应用修改')
    return True


# ── 子句表：属性语法的**全部**唯一写法都在这张表里 ─────────────────────────
#
# 解析器已收敛到 md_to_word.spec.parser（中英双语词表、封闭取值、1 对 1 纠错提示）。
# 这里只做转发：本文件负责「把 spec 写成 OOXML」，不负责「读懂格式要求文本」。
from md_to_word.spec.parser import (  # noqa: E402
    parse_requirements,
    parse_style_props,
    parse_table_props,
)

_S = "style"
_T = "table"
TABLE_ALIGN_MAP = {cn: v for cn, v in ALIGN_VAL.items() if cn != "两端对齐"}
TABLE_STYLE_MAP = {"三线表": "three-line", "全网格线": "grid", "顶底线": "top-bottom", "无边框": "none"}
TABLE_VALIGN_MAP = {"居中": "center", "顶端": "top", "底端": "bottom"}
TABLE_WIDTH_KEYS = {"线宽上": "top", "线宽底": "bottom", "线宽表头下": "header_bottom", "线宽内": "inside"}


def decor_token_to_field(tok):
    """把 {…} 标记翻成 Word 域指令；不认识返回 None（由调用方报错，不静默漏掉）"""
    if tok == '页码':
        return 'PAGE   \\* MERGEFORMAT '
    if tok == '总页数':
        return 'NUMPAGES   \\* MERGEFORMAT '
    m = _RE_STYLEREF.match(tok)
    if m:
        return 'STYLEREF  "%s"   \\* MERGEFORMAT ' % m.group(1)
    return None


def parse_decor_content(text, lineno=None, key=None):
    """把页眉/页脚内容串切成 [('text', s) | ('field', 域指令)]"""
    parts = []
    pos = 0
    for m in _RE_DECOR_TOKEN.finditer(text):
        if m.start() > pos:
            parts.append(('text', text[pos:m.start()]))
        tok = m.group(1).strip()
        instr = decor_token_to_field(tok)
        if instr is None:
            where = f'第 {lineno} 行' if lineno else ''
            whose = f'「{key}」' if key else ''
            raise ValueError(
                f'{where}{whose}未知的页眉/页脚域标记 "{{{tok}}}"；'
                f'可用：{{页码}} {{总页数}} {{样式标题:<样式名>}}')
        parts.append(('field', instr))
        pos = m.end()
    if pos < len(text):
        parts.append(('text', text[pos:]))
    return parts


@dataclass
class TemplateSpec:
    """requirements.txt 的**唯一**解析结果：各写入器都从这里取料"""
    styles: list = field(default_factory=list)
    default_latin: Optional[str] = None
    margins: Optional[dict] = None
    table: dict = field(default_factory=dict)          # 表格：对齐/边框/线宽/单元格垂直
    table_header: Optional[dict] = None                # 表头：表头行（firstRow）属性
    decor: dict = field(default_factory=dict)          # (header|footer, default|even|first) → 内容串
    page_number: Optional[tuple] = None                # (fmt, start)
    page_size: str = 'a4'                              # 纸张：a4|letter
    gutter_cm: Optional[float] = None                  # 装订线：左侧额外留出的装订距离
    doc_grid_lines: Optional[int] = None               # 文档网格：每页多少行
    proofing_lang: Optional[str] = None                # 校对语言（zh-CN 等）
    header_distance: Optional[float] = None            # 页眉距边界（cm）
    footer_distance: Optional[float] = None            # 页脚距边界（cm）


def _bad_value(lineno, key, desc, should, why=''):
    """行级键取值不认识时的**统一**报错：行号 + 键 + 原值 + 应写什么。"""
    return ValueError(
        f'第 {lineno} 行「{key}」的值不认识：{desc}'
        + (f'\n  {why}' if why else '')
        + f'\n  应写：{key}：{should}')


# ── 行级键（`<键>：<值>`）──
def _line_latin_font(spec, desc, lineno):
    spec.default_latin = desc


def _line_margins(spec, desc, lineno):
    spec.margins = parse_margins(desc, lineno)


def _line_proofing_lang(spec, desc, lineno):
    if not re.fullmatch(r'[A-Za-z]{2}(?:-[A-Za-z]{2,4})?', desc):
        raise _bad_value(lineno, '校对语言', desc, 'zh-CN', why='整行只写一个语言代码，不写语言名')
    spec.proofing_lang = desc


def _line_distance(field, key):
    """页眉/页脚距边界：单值 + 单位 cm。"""
    def handler(spec, desc, lineno):
        m = re.fullmatch(r'(\d+(?:\.\d+)?)cm', desc)
        if not m:
            raise _bad_value(lineno, key, desc, '1.4cm', why='整行只写一个值，单位用 cm（无空格）')
        setattr(spec, field, float(m.group(1)))
    return handler


def _line_gutter(spec, desc, lineno):
    """装订线：整行一个 cm 值。"""
    m = re.fullmatch(r'(\d+(?:\.\d+)?)cm', desc)
    if not m:
        raise _bad_value(lineno, '装订线', desc, '1cm', why='整行只写一个值，单位用 cm（无空格）')
    spec.gutter_cm = float(m.group(1))


def _line_doc_grid(spec, desc, lineno):
    """文档网格：`每页22行`。

    中文学位论文与公文常按「每页 N 行」规定版式（GB/T 9704 就是每页 22 行）。
    Word 里对应「页面设置 → 文档网格 → 指定行网格」，写的是 `w:docGrid` 的
    `linePitch`——每行占多少 twips，由版心高度除以行数算出。
    """
    m = re.fullmatch(r'每页(\d+)行', desc)
    if not m:
        raise _bad_value(lineno, '文档网格', desc, '每页22行',
                         why='只支持按行数指定；「每行多少字」得靠字号与字距自己凑')
    n = int(m.group(1))
    if not 1 <= n <= 200:
        raise _bad_value(lineno, '文档网格', desc, '每页22行', why='行数要在 1–200 之间')
    spec.doc_grid_lines = n


def _line_paper(spec, desc, lineno):
    if desc not in PAGE_SIZE_MAP:
        raise _bad_value(lineno, '纸张', desc, 'A4|Letter')
    spec.page_size = PAGE_SIZE_MAP[desc]


def _line_page_number(spec, desc, lineno):
    spec.page_number = parse_page_number(desc, lineno)


def _line_table(spec, desc, lineno):
    spec.table = parse_table_props(desc, lineno)


def _line_table_header(spec, desc, lineno):
    # 「跨页重复」由子句表消费（_check_clause_preconditions 里限定它只能出现在这一行）
    spec.table_header = parse_style_props('表头', desc, lineno)


def _line_decor(container, key):
    def handler(spec, desc, lineno):
        spec.decor[container] = [] if desc == '无' else parse_decor_content(desc, lineno, key)
    return handler


_LINE_KEYS = {
    '西文字体': _line_latin_font,
    '页边距': _line_margins,
    '校对语言': _line_proofing_lang,
    '页眉距边界': _line_distance('header_distance', '页眉距边界'),
    '页脚距边界': _line_distance('footer_distance', '页脚距边界'),
    '纸张': _line_paper,
    '装订线': _line_gutter,
    '文档网格': _line_doc_grid,
    '页码': _line_page_number,
    '表格': _line_table,
    '表头': _line_table_header,
    **{key: _line_decor(container, key) for key, container in DECOR_CONTENT_KEYS.items()},
}


def parse_page_number(desc, lineno=None):
    """解析「页码：<格式> [起始<N>]」→ (OOXML fmt, start)"""
    tokens = desc.split()
    if not (1 <= len(tokens) <= 2) or tokens[0] not in PAGE_NUMBER_FMT:
        raise _bad_value(lineno, '页码', desc, '阿拉伯 起始1',
                         why=f'第一个词须为 {"/".join(PAGE_NUMBER_FMT)} 之一，可选第二个词为 起始<正整数>')
    fmt = PAGE_NUMBER_FMT[tokens[0]]
    start = None
    if len(tokens) == 2:
        m = re.fullmatch(r'起始(\d+)', tokens[1])
        if not m:
            raise _bad_value(lineno, '页码', desc, '阿拉伯 起始1',
                             why=f'不认识「{tokens[1]}」，起始页码应写为 起始1')
        start = int(m.group(1))
    return fmt, start


_MARGIN_SIDE_MAP = {'上': 'top', '下': 'bottom', '左': 'left', '右': 'right'}
_MARGIN_TOKEN_RE = re.compile(r'^(?P<side>[上下左右])(?P<val>\d+(?:\.\d+)?)cm$')


def parse_margins(desc, lineno=None):
    """解析「页边距：上Xcm 下Xcm 左Xcm 右Xcm」。四边必填、单位只认 cm。"""
    margins = {}
    bad_tokens = []
    for tok in desc.split():
        m = _MARGIN_TOKEN_RE.fullmatch(tok)
        if not m:
            bad_tokens.append(tok)
            continue
        en = _MARGIN_SIDE_MAP[m.group('side')]
        if en in margins:
            raise _bad_value(lineno, '页边距', desc, '上2.54cm 下2.54cm 左2.54cm 右2.54cm',
                             why=f'「{m.group("side")}」重复写了两次')
        margins[en] = float(m.group('val'))
    missing = [cn for cn, en in _MARGIN_SIDE_MAP.items() if en not in margins]
    if missing or bad_tokens:
        why = []
        if missing:
            why.append(f'缺了「{"」「".join(missing)}」')
        if bad_tokens:
            why.append(f'不认识「{"」「".join(bad_tokens)}」')
        raise _bad_value(lineno, '页边距', desc, '上2.54cm 下2.54cm 左2.54cm 右2.54cm',
                         why='；'.join(why) + '（四边必填、单位用 cm，如 上2.54cm）')
    return margins


def apply_proofing_language(doc, lang):
    """把校对语言写进 docDefaults。

    只设 `eastAsia` 槽（中文、日文这些字符用的槽）——西文槽保持底模的 en-US，
    否则中文文档里的英文词会被拿去按中文词典查。底模两个槽都是 en-US，不写中文
    语言，Word 就用英文词典查中文，整篇标成拼写错误。
    """
    root = doc.styles.element
    defaults = root.find(W('docDefaults'))
    if defaults is None:
        defaults = etree.Element(W('docDefaults'))
        root.insert(0, defaults)
    rpr_default = ensure_ordered_child(defaults, 'rPrDefault', ['rPrDefault', 'pPrDefault'])
    rpr = ensure_ordered_child(rpr_default, 'rPr', ['rPr'])
    lang_el = ensure_ordered_child(rpr, 'lang', _RPR_ORDER)
    lang_el.set(W('eastAsia'), lang)


def get_default_template():
    """取纯净现代 Word 365 基底模板（resources/pure_modern_base.docx）。"""
    local_base = get_base_template_path()
    if not local_base.exists():
        raise TemplateError(
            f'未找到基础模板 {local_base}；请运行 scripts/create_pure_modern_base.py 重新生成。'
        )
    logger.info(f"使用内置纯净现代 Word 365 基底模板: {local_base}")
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.docx')
    shutil.copyfile(str(local_base), tmp.name)
    tmp.close()
    return tmp.name


def cleanup_temp(path: str):
    try:
        os.unlink(path)
    except OSError:
        pass


# ── 所有文档无条件必填的基础样式骨架 ──
# 无论导出普通文档/报告还是学术论文，标题、一级章节与正文三态（Body Text / First Paragraph / Normal）
# 都是必现基础段落；其余结构（二级及以上标题、摘要、关键词、图表、参考文献、目录、引用块）
# 均通过下方 STRUCT_TRIGGERS 按「稿件实际用到即必填」精确拦截。
REQUIRED_STYLES = {
    'Title': '文档标题',
    'Heading 1': '一级标题',
    'Body Text': '正文',
    'First Paragraph': '首段',
    'Normal': '普通',
}

# ── 条件成套必填（表格三件套）──
# 纯文字文档无表格时不强制写；一旦声明了「表格：」或「表头：」中的任意一项（或稿件中出现了表格），
# 则「表格」「表头」「紧凑（单元格文字）」必须成套配齐，防止表头或表内文字静默回退默认样式。
REQUIRED_ROWS = {}

# ── **条件必填**：写了内容就必须给样式行 ──
# 页眉/页脚本身可有可无（不写＝不加，不算错）；但写了「页眉内容：」却不给「页眉：」
# 样式行，页眉就会套用底模默认字体——正是要防的那种静默错误。
CONDITIONAL_STYLE_ROWS = {
    'Header': ('页眉', '页眉：宋体/Times New Roman 小五 不加粗 不倾斜 居中 行距1.0倍 段前0磅 段后0磅 首行缩进0字符'),
    'Footer': ('页脚', '页脚：宋体/Times New Roman 小五 不加粗 不倾斜 居中 行距1.0倍 段前0磅 段后0磅 首行缩进0字符'),
}

# ── **条件成套必填（学位论文目录）** ──
# 无目录时不强制写；一旦声明了目录相关样式中的任意一个，
# 目录标题 + 一/二/三级目录条目必须成套写齐，防止漏写某一级导致该级退回 Word 默认缩进。
TOC_STYLE_GROUP = {
    'TOC Heading': '目录标题',
    'TOC 1': '目录一级',
    'TOC 2': '目录二级',
    'TOC 3': '目录三级',
}

# ── **稿件结构 → 必须声明的样式** ──
# 与 REQUIRED_STYLES（基础骨架无条件必填）互补：普通文档未必有摘要/关键词/参考文献/图表，
# 不必强求纯文档也声明全套学术论文样式；但**稿件里真出现了**这些结构，要求文件就必须
# 显式定义，否则静默回退到底模默认样式。
STRUCT_TRIGGERS = [
    (re.compile(r'^##\s+\S'), '二级标题 (##)', [('二级标题', 'Heading 2')]),
    (re.compile(r'^###\s+\S'), '三级标题 (###)', [('三级标题', 'Heading 3')]),
    (re.compile(r'^####\s+\S'), '四级标题 (####)', [('四级标题', 'Heading 4')]),
    (re.compile(r'^#####\s+\S'), '五级标题 (#####)', [('五级标题', 'Heading 5')]),
    (re.compile(r'^######\s+\S'), '六级标题 (######)', [('六级标题', 'Heading 6')]),
    (TOC_MARKER_RE, '目录短代码 ({{< toc >}})',
     [(cn, en) for en, cn in TOC_STYLE_GROUP.items()]),
    (re.compile(r'^>\s+\S'), '引用块 (>)', [('引用块', 'Block Text')]),
    (re.compile(r'custom-style=(["\'])Abstract Title\1'), '摘要标题包裹', [('摘要标题', 'Abstract Title')]),
    (re.compile(r'custom-style=(["\'])Abstract\1'), '摘要正文包裹', [('摘要', 'Abstract')]),
    (re.compile(r'custom-style=(["\'])Keywords\1'), '关键词包裹', [('关键词', 'Keywords')]),
    (re.compile(r'custom-style=(["\'])References Title\1'), '参考文献标题包裹', [('参考文献标题', 'References Title')]),
    (re.compile(r'(?:::\s*\{#refs\}|custom-style=(["\'])Bibliography\1)'), '参考文献区块', [('参考文献', 'Bibliography')]),
    (re.compile(r'!\[[^\]]*\]\([^)]+\)'), '图片 (![](...))', [('图', 'Captioned Figure'), ('图片题注', 'Image Caption')]),
    (re.compile(r'custom-style=(["\'])Table Caption\1'), '表格题注包裹', [('表格题注', 'Table Caption')]),
    (re.compile(r'custom-style=(["\'])Table Note\1'), '表注包裹', [('表注', 'Table Note')]),
    (re.compile(r'^\s*\|.*\|'), 'Markdown 表格 (|...|)', [
        ('紧凑', 'Compact'),
        ('表格', '__ROW_table__'),
        ('表头', '__ROW_table_header__'),
    ]),
]


def missing_styles_for_body(masked_body: str, declared: set, start_line: int = 1) -> list:
    """稿件正文用到的结构 → 要求文件漏声明的样式或配置行。

    `masked_body` 是**已掩码代码块**的正文（`skill_utils.mask_code_fences`）——方法学论文
    常在代码块里演示语法，那里的 `>`、`###` 不是稿件结构。
    返回 [(行号, 结构描述, 中文写法, Word 样式名)]，空列表 = 都声明了。
    """
    issues, reported = [], set()
    for i, line in enumerate(masked_body.split('\n'), start_line):
        for pat, feat, styles in STRUCT_TRIGGERS:
            if not pat.search(line):
                continue
            for cn, en in styles:
                if en not in declared and (feat, en) not in reported:
                    reported.add((feat, en))
                    issues.append((i, feat, cn, en))
    return issues

_PARA = ['font_name', 'font_size', 'alignment', 'line_spacing', 'space_before', 'space_after']
# 「字距」是可选的字符级属性，**不进下面的必填表**——列进去就等于逼每个样式行都写它，
# 现有格式要求文件会全部判为缺字段。
_TITLE = _PARA + ['bold', 'italic']   # 标题类：加粗/倾斜状态必须显式，否则继承底模设定

REQUIRED_FIELDS_BY_STYLE = {
    k: v + ['indent'] for k, v in {
        'Title': _TITLE,
        'Subtitle': _TITLE,
        'Heading 1': _TITLE,
        'Heading 2': _TITLE,
        'Heading 3': _TITLE,
        'Heading 4': _TITLE,
        'Heading 5': _TITLE,
        'Heading 6': _TITLE,
        'Abstract Title': _TITLE,
        'References Title': _TITLE,
        'TOC Heading': _TITLE,
        # 目录条目同样是 basedOn Normal，不显式写缩进就会被正文的 2 字符带跑
        'TOC 1': _PARA,
        'TOC 2': _PARA,
        'TOC 3': _PARA,
        'Abstract': _PARA,
        'Keywords': _PARA,
        'Table Caption': _PARA,
        'Table Note': _PARA,
        'Image Caption': _PARA,
        'Body Text': _PARA,
        'First Paragraph': _PARA,
        'Compact': _PARA,
        'Normal': _PARA,
        'Bibliography': _PARA,
        'Block Text': _PARA,
        'Footnote Text': _PARA,
        'Author': _PARA,
        'Date': _PARA,
        'Header': _PARA,
        'Footer': _PARA,
        # 图片段落不带文字，字体字号无意义；要紧的只有缩进 0 与居中
        'Captioned Figure': ['alignment', 'line_spacing', 'space_before', 'space_after'],
    }.items()
}
# 「缩进」为什么每个契约样式都必填：不是每个样式都需要缩进，而是**不写就会继承**。
# 正文写「首行缩进2字符」之后 Normal 就带上了 firstLineChars=200，凡 basedOn Normal
# 又没自己写 w:ind 的样式都会跟着缩进 2 字符。图片段落、表格单元格这类一行就结束的，
# 被顶右 2 字符再叠加居中会顶出右边界——肉眼看是"图片没居中"，实则是继承来的缩进。

FIELD_CN = {
    'font_name': '字体',
    'font_size': '字号',
    'alignment': '对齐方式',
    'line_spacing': '行距',
    'space_before': '段前',
    'space_after': '段后',
    'indent': '缩进（首行缩进/悬挂缩进，不需则填0）',
    'bold': '加粗状态',
    'italic': '倾斜状态',
}

FIELD_FIX_HINT = {
    'font_name': '字体（如"宋体"或"Times New Roman"或"宋体/Times New Roman"）',
    'font_size': '字号（如"12pt"或"小四"）',
    'alignment': '对齐方式（如"居中"/"左对齐"/"右对齐"/"两端对齐"）',
    'line_spacing': '行距（"行距2倍" / "行距20磅" / "行距最小值25磅"，不能用0）',
    'space_before': '段前（如"段前0磅"，不需段间距填0）',
    'space_after': '段后（如"段后6磅"，不需段间距填0）',
    'indent': '缩进（正文写"首行缩进2字符"，参考文献写"悬挂缩进0.5英寸"，表格写"首行缩进0字符"）',
    'bold': '加粗状态（如"加粗"或"不加粗"，标题类样式必须显式指定）',
    'italic': '倾斜状态（如"倾斜"或"不倾斜"，标题类样式必须显式指定）',
}


def _check_required_fields(spec) -> bool:
    """校验必填项是否齐备：必填样式 / 样式的必填字段 / 条件成套必填行。"""
    req_by_style = {r['style_name']: r for r in spec.styles}

    missing_required = []   # (别名, 字段)
    missing_styles = []     # (样式名, 别名)
    missing_rows = []       # (行名, 应写)

    for style_name, alias in REQUIRED_STYLES.items():
        if style_name not in req_by_style:
            missing_styles.append((style_name, alias))
            continue
        req = req_by_style[style_name]
        required_fields = REQUIRED_FIELDS_BY_STYLE.get(style_name, [])
        for fname in required_fields:
            if fname == 'indent':
                if 'indent' not in req or 'indent_type' not in req:
                    missing_required.append((alias, fname))
            elif fname not in req or req[fname] is None or req[fname] == '':
                missing_required.append((alias, fname))

    # 检查用户声明了但非无条件必选的样式（如二~六级标题、摘要、参考文献、页眉页脚、目录等）：声明了就必须全字段显式覆盖
    for req in spec.styles:
        style_name = req['style_name']
        alias = req.get('alias') or STYLE_CN.get(style_name, style_name)
        if style_name in REQUIRED_FIELDS_BY_STYLE and style_name not in REQUIRED_STYLES:
            required_fields = REQUIRED_FIELDS_BY_STYLE[style_name]
            for fname in required_fields:
                if fname == 'indent':
                    if 'indent' not in req or 'indent_type' not in req:
                        missing_required.append((alias, fname))
                elif fname not in req or req[fname] is None or req[fname] == '':
                    missing_required.append((alias, fname))

    # 条件成套必填 1：若声明了目录样式组中的任意一个，必须成套写齐全部 4 项
    if any(s in req_by_style for s in TOC_STYLE_GROUP):
        for s_name, s_alias in TOC_STYLE_GROUP.items():
            if s_name not in req_by_style:
                missing_styles.append((s_name, f'{s_alias}（已声明目录样式，须成套配齐）'))

    # 条件成套必填 2：若声明了「表格：」或「表头：」，则表格、表头、紧凑（单元格文字）三者必须配齐
    if spec.table or spec.table_header:
        if not spec.table:
            missing_rows.append(('表格', '表格：居中 三线表 线宽上1磅 线宽底1磅 线宽表头下0.25磅 单元格垂直居中'))
        if not spec.table_header:
            missing_rows.append(('表头', '表头：宋体/Times New Roman 五号 加粗 不倾斜 居中 行距1.0倍 段前0磅 段后0磅 首行缩进0字符'))
        if 'Compact' not in req_by_style:
            missing_styles.append(('Compact', '紧凑（已声明表格样式，须同时声明表内单元格文字样式）'))

    # 条件成套必填 3：图片与图片题注成套配齐
    if ('Captioned Figure' in req_by_style) != ('Image Caption' in req_by_style):
        if 'Captioned Figure' not in req_by_style:
            missing_styles.append(('Captioned Figure', '图（已声明图片题注，须同时声明图片段落样式）'))
        else:
            missing_styles.append(('Image Caption', '图片题注（已声明图样式，须同时声明图片题注样式）'))

    # 条件成套必填 4：写了摘要标题须配摘要正文；写了参考文献标题须配参考文献条目
    if 'Abstract Title' in req_by_style and 'Abstract' not in req_by_style:
        missing_styles.append(('Abstract', '摘要（已声明摘要标题，须同时声明摘要正文样式）'))
    if 'References Title' in req_by_style and 'Bibliography' not in req_by_style:
        missing_styles.append(('Bibliography', '参考文献（已声明参考文献标题，须同时声明参考文献条目样式）'))

    for key, (name, example) in REQUIRED_ROWS.items():
        if not getattr(spec, key, None):
            missing_rows.append((name, example))

    # 表头行（firstRow）字段完整性校验
    if spec.table_header:
        for fname in _PARA + ['indent']:
            if fname == 'indent':
                if 'indent' not in spec.table_header or 'indent_type' not in spec.table_header:
                    missing_required.append(('表头', fname))
            elif (fname not in spec.table_header or spec.table_header[fname] is None
                  or spec.table_header[fname] == ''):
                missing_required.append(('表头', fname))

    for style_name, (name, example) in CONDITIONAL_STYLE_ROWS.items():
        written = any(kind == ('header' if style_name == 'Header' else 'footer')
                      for kind, _ in spec.decor)
        if written and style_name not in req_by_style:
            missing_rows.append((name, example))

    has_errors = bool(missing_required or missing_styles or missing_rows)

    if missing_styles:
        logger.info('\n⚠️ 必填样式缺失（将沿用 Word 默认值，可能不符合期刊要求）:')
        for style_name, alias in missing_styles:
            logger.info(f' - {alias}（{style_name}）')
        logger.info(' 建议在 requirements.txt 中补充这些样式的定义')

    if missing_rows:
        logger.info('\n⚠️ 必填行缺失（不写就会套用默认样式）:')
        for name, example in missing_rows:
            logger.info(f' - {name}: 应写「{example}」')

    if missing_required:
        logger.info('\n⚠️ 必填字段缺失（缺失会沿用 Word 默认值，可能导致格式不合规）:')
        for alias, field in missing_required:
            hint = FIELD_FIX_HINT.get(field, field)
            logger.info(f' - {alias}: 缺少{FIELD_CN.get(field, field)}')
            logger.info(f' 应写: {hint}')

    if has_errors:
        logger.error('\n❌ 必填项或必填字段缺失，终止生成。')
        logger.error(' 请按上方提示补齐缺失项（不需要间距或缩进时，也须显式填 0，如「段前0磅」「首行缩进0字符」）。')
    else:
        logger.info('✅ 必填字段校验通过')

    return has_errors


# ══════════════════════════════════════════════════════════════════
#  页眉 / 页脚 / 页码
#
#  Pandoc 的 --reference-doc **会**连同参考文档的 header/footer 部件与 sectPr 一起
#  带进产物（官方手册明写 "document properties (including margins, page size,
#  header, and footer) are used"）。但 pandoc 只产出**一个** section——取参考文档
#  的第一个 sectPr。因此本技能把整份文档做成单节，奇偶页差异走 evenAndOddHeaders
#  + default/even 两个 headerReference，而不是靠多节。
# ══════════════════════════════════════════════════════════════════

_XML_SPACE = '{http://www.w3.org/XML/1998/namespace}space'


def _clear_runs(paragraph):
    for r in list(paragraph.runs):
        r._r.getparent().remove(r._r)


def append_field(paragraph, instr):
    """在段落里插入一个 Word 域（begin → instrText → separate → end）。

    不写缓存结果：值由 Word 在更新域/打印时算出，故导出链在出 PDF 前会
    调用 Word 的 update field（页眉里的 STYLEREF 尤其需要）。
    """
    r = paragraph.add_run()._r
    etree.SubElement(r, W('fldChar')).set(W('fldCharType'), 'begin')
    it = etree.SubElement(r, W('instrText'))
    it.set(_XML_SPACE, 'preserve')
    it.text = instr
    etree.SubElement(r, W('fldChar')).set(W('fldCharType'), 'separate')
    etree.SubElement(r, W('fldChar')).set(W('fldCharType'), 'end')


def write_decor_container(container, parts, style_name):
    """把 parse_decor_content() 的结果写进一个页眉/页脚容器"""
    container.is_linked_to_previous = False   # 先强制建出部件，否则写进去会被继承覆盖
    p = container.paragraphs[0] if container.paragraphs else container.add_paragraph()
    _clear_runs(p)
    p.style = style_name
    for kind, val in parts:
        if kind == 'text':
            p.add_run(val)
        else:
            append_field(p, val)


def collect_styleref_names(spec):
    """取出所有 {样式标题:X} 里的 X（去重、保序）"""
    names = []
    for parts in spec.decor.values():
        for kind, val in parts:
            if kind != 'field':
                continue
            m = re.match(r'STYLEREF\s+"(.*?)"', val)
            if m and m.group(1) not in names:
                names.append(m.group(1))
    return names


def ensure_style_alias(doc, style_name, alias):
    """把 alias 注册为样式的 w:aliases，返回样式是否存在于模板。

    **为什么需要**：Word 的 STYLEREF 域按样式的**本地化显示名或别名**匹配，唯独不认
    w:name 的字面值。实测（中文版 Word）：样式 w:name="heading 1" 时，
    STYLEREF "heading 1" 与 STYLEREF "Heading 1" 都报"错误!未找到引用源"，
    只有 STYLEREF "标题 1" 能解析；给该样式加上别名 "Heading 1" 后两者才都能用。
    因此这里把要求在文件里写的那个名字原样注册成别名——用户写什么，域就用什么，
    不依赖使用者 Word 的界面语言。
    """
    try:
        style = doc.styles[style_name]
    except KeyError:
        return False

    el = style.element
    al = el.find(W('aliases'))
    existing = [a.strip() for a in (al.get(W('val')) or '').split(',') if a.strip()] if al is not None else []
    if alias in existing:
        return True
    existing.append(alias)
    if al is None:
        al = etree.Element(W('aliases'))
        name_el = el.find(W('name'))
        if name_el is not None:
            name_el.addnext(al)
        else:
            el.insert(0, al)
    al.set(W('val'), ','.join(existing))
    return True


def apply_page_decorations(doc, spec):
    """写入页眉/页脚内容、奇偶与首页开关、页码格式。

    只在 requirements 里**真的出现了**的容器上动手：没写的容器保持"链接上一节"，
    不会凭空建出空页眉把偶数页刷白。
    """
    if not spec.decor and not spec.page_number:
        return

    if len(doc.sections) != 1:
        raise TemplateError(
            f'模板必须是单节（当前 {len(doc.sections)} 节），多节模板的页眉页脚会错位。'
        )

    section = doc.sections[0]

    uses_even = any(which == 'even' for _, which in spec.decor)
    uses_first = any(which == 'first' for _, which in spec.decor)

    # 一旦开启奇偶/首页不同，**没显式写的那几套容器会变成空白，而不是继承默认那一套**。
    # 例如只写「页脚内容：{页码}」再开奇偶不同，偶数页会整页没有页码。
    # 故把默认内容补进空缺的槽位——让「页码每页都有」这类常理成立；要覆盖就显式写。
    for kind, cn in (('header', '页眉'), ('footer', '页脚')):
        base = spec.decor.get((kind, 'default'))
        if base is None:
            continue
        for which, label in (('even', '偶数页'), ('first', '首页')):
            if (uses_even if which == 'even' else uses_first) and (kind, which) not in spec.decor:
                spec.decor[(kind, which)] = base
                logger.info(f' ✓ {label}{cn}: 未单独指定，已沿用默认{cn}内容')

    if uses_even:
        doc.settings.odd_and_even_pages_header_footer = True
        logger.info(' ✓ 页眉页脚: 已开启奇偶页不同')
    if uses_first:
        section.different_first_page_header_footer = True
        logger.info(' ✓ 页眉页脚: 已开启首页不同')

    which_prefix = {'default': '', 'even': 'even_page_', 'first': 'first_page_'}
    kind_name = {'header': ('header', 'Header', '页眉'), 'footer': ('footer', 'Footer', '页脚')}
    for (kind, which), parts in spec.decor.items():
        attr, style_name, cn = kind_name[kind]
        write_decor_container(getattr(section, which_prefix[which] + attr), parts, style_name)
        label = {'default': '', 'even': '偶数页', 'first': '首页'}[which]
        logger.info(f' ✓ {label}{cn}: 已设置内容与域')

    for nm in collect_styleref_names(spec):
        if ensure_style_alias(doc, nm, nm):
            logger.info(f' ✓ 页眉/页脚域: 已把「{nm}」注册为样式别名（STYLEREF 才能解析）')
            continue
        # STYLEREF 只认 **Word 样式名**（Heading 1 这种），中文写法（一级标题）不是样式名，
        # 写进去 Word 会在页眉显示「错误!未找到引用源」。这里是导出最后一道能拦的地方。
        cn = STYLE_MAP.get(nm)
        hint = (f'\n  「{nm}」是本技能的中文写法，对应 Word 样式名「{cn}」——'
                f'域里要写后者：{{样式标题:{cn}}}') if cn else ''
        raise TemplateError(
            f'页眉/页脚域 {{样式标题:{nm}}} 指向的样式在模板里不存在。'
            f'STYLEREF 只认 Word 样式名（如 Heading 1）。{hint}'
        )

    # 「页码：」只往 sectPr 写编号格式与起始值；页码要真的印在纸上，得有一行页眉/页脚
    # 内容带 {页码} 域。写了「页码：」却没有 {页码} 域 = 写了等于没写，且三道闸门全绿。
    if spec.page_number and not any(
            kind == 'field' and val.split()[0] == 'PAGE'
            for parts in spec.decor.values() for kind, val in parts):
        raise TemplateError(
            '写了「页码：」，但没有任何页眉/页脚内容带 {页码} 域。'
            '「页码：」只定格式与起始值，把页码画到纸上的是 {页码} 域；'
            '补一行「页脚内容：{页码}」，或删掉「页码：」这一行。'
        )

    if spec.page_number:
        fmt, start = spec.page_number
        pg = ensure_ordered_child(section._sectPr, 'pgNumType', _SECTPR_ORDER)
        pg.set(W('fmt'), fmt)
        if start is not None:
            pg.set(W('start'), str(start))
        else:
            pg.attrib.pop(W('start'), None)
        logger.info(f' ✓ 页码: 格式={fmt}' + (f' 起始={start}' if start is not None else ''))


# ══════════════════════════════════════════════════════════════════
#  表格
#
#  Pandoc 把**表格内所有段落**映射为 Compact 样式（实测：Compact 在表内 582 次、
#  表外 0 次），所以表内文字的字体/字号/缩进由「紧凑：」样式行控制；而**表头行**
#  与**单元格**无法用段落样式区分，只能写进 Table 样式的条件格式
#  （w:tblStylePr type="firstRow" / "wholeTable"）。
#
#  ⚠ 条件格式里**边框走 tcPr、字符/段落属性走 rPr/pPr**——容器写错的那部分会
#    **静默不生效**，而模板 XML 看上去完全正常。详见 _write_cell_borders 的 docstring。
# ══════════════════════════════════════════════════════════════════

# 各边框样式的默认线宽（磅）——与历史行为一致，未显式声明时用它
TABLE_DEFAULT_WIDTHS = {
    'three-line': {'top': 1.5, 'bottom': 1.5, 'header_bottom': 0.75},
    'grid': {'inside': 0.5},
    'top-bottom': {'top': 1.5, 'bottom': 1.5},
}

TABLE_BORDER_SIDES = ('top', 'left', 'bottom', 'right', 'insideH', 'insideV')

# ── 表格单元格左右内边距：**声明项**，不写即 0 ──
# 要求文件写「单元格内边距0.19cm」即生效；不写时用 0，且**显式写进产物**——
# 不能让底模的 Table→Normal Table 继承值替我们决定（继承来的值没人声明过，却实实在在
# 决定产物长相）。默认 0 不是审美选择：渲染器把 `tblW` 写成 100%、`tblLayout` 写成 fixed，
# 列宽由 Markdown 分隔行的比例决定，**列宽之和恒等于版心宽度**——内边距没有"额外空间"
# 可占，只能从列里扣。默认取 0，"分隔行的比例"才等于"文字实际能用的宽度"；要留白就写出来，
# 那时少掉的那点宽度是声明过的、可追责的。
TABLE_CELL_MARGIN_DEFAULT_CM = 0.0


def table_border_widths(style: str, declared: dict) -> dict:
    """表格样式 + 要求文件声明的线宽 → {边: 磅值，None = 不画线}。

    **生成（写 XML）与导出后对账（读 XML）都从这里取期望值**：两边各推一次时，
    「全网格线的上下外框能不能单独加粗」这种细节会两处不一致——对账按 `线宽内` 取期望，
    于是把"写了 `线宽上` 却不生效"的产物判成通过。
    """
    w = dict(TABLE_DEFAULT_WIDTHS.get(style, {}))
    w.update(declared or {})
    if style == 'three-line':
        return {'top': w.get('top', 1.5), 'bottom': w.get('bottom', 1.5),
                'left': None, 'right': None, 'insideH': None, 'insideV': None}
    if style == 'top-bottom':
        return {side: (w.get(side, 1.5) if side in ('top', 'bottom') else None)
                for side in TABLE_BORDER_SIDES}
    if style == 'grid':
        # 全网格线：上下外框可单独加粗（`线宽上`/`线宽底`），左右与内部横竖线用 `线宽内`
        line = w.get('inside', 0.5)
        return {'top': w.get('top', line), 'bottom': w.get('bottom', line),
                'left': line, 'right': line, 'insideH': line, 'insideV': line}
    return {}


def _pt_to_border_sz(pt):
    """磅 → OOXML 边框 sz（1/8 磅），下限 2（1/4 磅，Word 最细单实线）"""
    return str(max(2, int(round(pt * 8))))


def _get_tbl_style_pr(style_element, cond_type):
    """取/建 w:tblStylePr[type=…]。同类型只能有一个，边框与表头格式共用它。"""
    for el in style_element.findall(W('tblStylePr')):
        if el.get(W('type')) == cond_type:
            return el
    el = etree.SubElement(style_element, W('tblStylePr'))
    el.set(W('type'), cond_type)
    return el


def _add_border(parent, side, sz='none'):
    el = etree.SubElement(parent, W(side))
    if sz == 'none':
        el.set(W('val'), 'none')
    else:
        el.set(W('val'), 'single')
        el.set(W('sz'), sz)
        el.set(W('space'), '0')
        el.set(W('color'), '000000')


def _write_borders(tbl_pr, sides):
    """把 w:tblBorders 重建到 tbl_pr（幂等）。

    `sides` 是 {边: 线宽磅|'none'}，**只写列出的边**。
    条件格式（tblStylePr）里的 tblBorders 是**按边覆盖**的：没列出的边会继承
    全表边框；若在此处补一圈 'none'，反而会把三线表的顶线/底线抹掉。
    """
    old = tbl_pr.find(W('tblBorders'))
    if old is not None:
        tbl_pr.remove(old)
    borders = ensure_ordered_child(tbl_pr, 'tblBorders', _TBLPR_ORDER)
    for side, width in sides.items():
        _add_border(borders, side, 'none' if width is None else _pt_to_border_sz(width))
    return borders


def _write_cell_borders(tbl_style_pr, sides):
    """把边框写进条件格式的 `tcPr/tcBorders`（**单元格级**）。

    ⚠ 条件格式（tblStylePr）里的边框**只能**走这里，不能走 `tblPr/tblBorders`：
    写在表级会被全表 `tblBorders` 的 `insideH="none"` 压掉，Word 渲染时那条线
    根本不出现——而 XML 看上去完全正常。实测三线表的表头下线：写表级 → 只有
    上下两条线；写单元格级 → 三条线齐。
    （同一个条件格式里的 `pPr`/`rPr`——对齐、字体——一直是正常生效的，
      所以只改边框这一处、不逐条渲染核对，极易漏判成"三线表没问题"。）
    """
    tc_pr = ensure_ordered_child(tbl_style_pr, 'tcPr', _TBLSTYLEPR_ORDER)
    old = tc_pr.find(W('tcBorders'))
    if old is not None:
        tc_pr.remove(old)
    borders = ensure_ordered_child(tc_pr, 'tcBorders', _TCPR_ORDER)
    for side, width in sides.items():
        _add_border(borders, side, 'none' if width is None else _pt_to_border_sz(width))
    return borders


def apply_table_format(doc, spec, default_latin=None):
    """写 Table 样式：对齐、边框（含三线表线宽）、表头行条件格式、单元格垂直对齐"""
    table = spec.table or {}
    table_align = table.get('align', 'center')
    table_style = table.get('style') or 'three-line'
    if table_style not in ('three-line', 'grid', 'top-bottom', 'none'):
        raise TemplateError(
            f'不支持的表格样式 "{table_style}"；可用：三线表 / 全网格线 / 顶底线 / 无边框。'
        )
    # 线宽旋钮只在**画得出来**的地方有意义：声明了却不生效比报错难查得多。
    declared_widths = table.get('widths') or {}
    if declared_widths and table_style == 'none':
        raise TemplateError(
            f'表格样式为「无边框」，不画任何线，不能写线宽（收到 {declared_widths}）；'
            f'要保留某条线请改用「三线表」或「顶底线」。'
        )
    # 「线宽内」只对全网格线有意义：三线表/顶底线按定义没有内部横线，写了也不会生效。
    if 'inside' in declared_widths and table_style in ('three-line', 'top-bottom'):
        raise TemplateError(
            f'表格样式为「{table_style}」，没有内部横线，不能写「线宽内」；'
            f'三线表的表头下线请用「线宽表头下」。'
        )

    try:
        style_el = doc.styles['Table'].element
    except KeyError:
        logger.info(' ⚠ 表格样式"Table"不存在（跳过表格边框和对齐设置）')
        return

    tbl_pr = ensure_ordered_child(style_el, 'tblPr', _STYLE_ORDER)

    # jc 必须排在 tblBorders 之前，否则 Word 忽略 jc（表格仍左对齐）
    old_jc = tbl_pr.find(W('jc'))
    if old_jc is not None:
        tbl_pr.remove(old_jc)
    jc_el = ensure_ordered_child(tbl_pr, 'jc', _TBLPR_ORDER)
    jc_el.set(W('val'), table_align)
    logger.info(f' ✓ Table: 表格对齐={table_align}（通过 Table 样式 jc 控制）')

    if table_style != 'none':
        declared = table.get('widths', {})
        widths = dict(TABLE_DEFAULT_WIDTHS.get(table_style, {}))
        widths.update(declared)                    # requirements 显式声明的线宽覆盖默认
        style_names = {'three-line': '三线表', 'grid': '全网格线', 'top-bottom': '顶底线'}
        logger.info(f' ✓ Table: 已应用{style_names.get(table_style, table_style)}边框'
              + (f'（线宽 {widths} 磅）' if declared else ''))
        _write_borders(tbl_pr, table_border_widths(table_style, declared))

        # 表头下线属「表头行」条件区域，必须写到**单元格级**（表级会被 insideH 压掉）。
        # 三线表按定义就该有这条线（默认 0.75 磅）；全网格线只有显式声明才写，
        # 否则表头那条线会跟周围网格线不一样粗。
        if table_style == 'three-line' or 'header_bottom' in declared:
            _write_cell_borders(_get_tbl_style_pr(style_el, 'firstRow'),
                                {'bottom': widths.get('header_bottom', 0.75)})

    # ── 表头行格式（复用段落的属性写入器，只是容器换成 tblStylePr）──
    if spec.table_header:
        props = spec.table_header
        cjk, latin = split_fonts(props.get('font_name'), default_latin)
        tsp = _get_tbl_style_pr(style_el, 'firstRow')
        write_run_props(ensure_ordered_child(tsp, 'rPr', _TBLSTYLEPR_ORDER), props, cjk=cjk, latin=latin)
        write_para_props(ensure_ordered_child(tsp, 'pPr', _TBLSTYLEPR_ORDER), props)
        if props.get('repeat_header'):
            tr_pr = ensure_ordered_child(tsp, 'trPr', _TBLSTYLEPR_ORDER)
            ensure_ordered_child(tr_pr, 'tblHeader', _TRPR_ORDER)
        logger.info(' ✓ Table: 已设置表头行格式'
              + ('（含跨页重复标题行）' if props.get('repeat_header') else ''))

    # 单元格内边距：要求文件写「单元格内边距Ncm」才生效，不写用 0。四种表格样式都适用。
    # 无论声明与否都要**写进产物**：底模的 Table→Normal Table 继承链里有 0.19cm，
    # 不写就等于默认被它决定（见 TABLE_CELL_MARGIN_DEFAULT_CM）。
    margin_twips = int(round(table.get('cell_margin_cm', TABLE_CELL_MARGIN_DEFAULT_CM) * _TWIPS_PER_CM))
    cell_mar = ensure_ordered_child(tbl_pr, 'tblCellMar', _TBLPR_ORDER)
    for side, twips in (('top', 0), ('left', margin_twips),
                        ('bottom', 0), ('right', margin_twips)):
        el = ensure_ordered_child(cell_mar, side, ('top', 'left', 'bottom', 'right'))
        el.set(W('w'), str(twips))
        el.set(W('type'), 'dxa')
    if 'cell_margin_cm' in table:
        logger.info(f' ✓ Table: 单元格内边距={table["cell_margin_cm"]}cm')

    if table.get('valign'):
        tsp = _get_tbl_style_pr(style_el, 'wholeTable')
        tc_pr = ensure_ordered_child(tsp, 'tcPr', _TBLSTYLEPR_ORDER)
        ensure_ordered_child(tc_pr, 'vAlign', _TCPR_ORDER).set(W('val'), table['valign'])
        logger.info(f' ✓ Table: 单元格垂直对齐={table["valign"]}')


def _props_from_req(req, page_size, margins):
    """样式行解析结果 → modify_style 吃的那份属性 dict。

    唯一一份：主循环与脚注默认都走这里，免得两条路上的可选属性各自漏一项。
    """
    props = {
        'font_name': req.get('font_name'),
        'font_size': req.get('font_size'),
        'bold': req.get('bold', False),
        'italic': req.get('italic', False),
        'alignment': req.get('alignment'),
        'line_spacing': req.get('line_spacing'),
        'indent': req.get('indent'),
        'indent_type': req.get('indent_type'),
        'indent_unit': req.get('indent_unit', 'pt'),
    }
    for key in ('space_before', 'space_after', 'no_border', 'bottom_border',
                'border_width_pt', 'left_indent_chars', 'char_spacing_pt',
                'keep_lines', 'page_break_before', 'widow_control'):
        if key in req:
            props[key] = req[key]
    if req.get('dot_leader_tab'):
        props['dot_leader_tab'] = True
        props['dot_leader_pos'] = _content_width_twips(page_size, margins)
    if req.get('join_toc'):
        props['join_toc'] = True
    if req.get('keep_next'):
        props['keep_next'] = True
    return props


# ── 脚注的两个样式 ──
# 脚注段落与脚注序号是渲染器**一定会打标记**的结构（`FootnoteText` / `FootnoteReference`），
# 而标记指不到定义时，OOXML 里不是"无格式"，是落到默认段落样式 Normal——也就是正文那一套
# （实测：脚注与正文同为 12pt、序号不是上标）。所以这两个样式**写没写要求文件都要落定义**。
_FOOTNOTE_STYLE = 'Footnote Text'
_FOOTNOTE_REF_STYLE = 'Footnote Reference'
_ALIGN_CN = {val: cn for cn, val in ALIGN_VAL.items()}


def _footnote_default_line(spec) -> str:
    """没写「脚注：」时的默认样式行：字体与对齐跟着正文走，字号按 SIZE_MAP 降一档。

    降一档的依据是本技能自己的排版原则（注释比正文小 1.5~2 磅）：小四正文 → 五号脚注。
    行距单倍、段前后 0、首行缩进 0——脚注是几行的注释，跟正文同为两端对齐时再带
    2 字符首行缩进会顶成锯齿。
    """
    body = next((r for r in spec.styles if r.get('style_name') == 'Body Text'), None) or {}
    fonts = '/'.join(body.get('font_name') or ['宋体', 'Times New Roman'])
    align = _ALIGN_CN.get(body.get('alignment') or '', '两端对齐')
    size_pt = (body.get('font_size') or SIZE_MAP['小四']).pt
    names = list(SIZE_MAP)
    i = next((n for n, name in enumerate(names) if SIZE_MAP[name].pt <= size_pt), len(names) - 1)
    size_name = names[min(i + 1, len(names) - 1)]
    return (f'{fonts} {size_name} 不加粗 不倾斜 {align} '
            f'行距1.0倍 段前0磅 段后0磅 首行缩进0字符')


def apply_footnote_format(doc, spec, page_size, margins, default_latin=None):
    """给两个脚注样式落定义：写了「脚注：」用写的，没写从正文推一套默认。

    另：脚注序号唯一要紧的属性是**上标**——不设它，正文里的序号就是个大号数字，
    读起来和普通数字分不开。
    """
    declared = any(r.get('style_name') == _FOOTNOTE_STYLE for r in spec.styles)
    if not declared:
        req = parse_style_props(_FOOTNOTE_STYLE, _footnote_default_line(spec))
        modify_style(doc, _FOOTNOTE_STYLE,
                     _props_from_req(req, page_size, margins), default_latin=default_latin)
        logger.info(f' ✓ {_FOOTNOTE_STYLE}: 要求文件未写「脚注」，已按正文推定默认')

    try:
        ref = doc.styles[_FOOTNOTE_REF_STYLE]
    except KeyError:
        ref = doc.styles.add_style(_FOOTNOTE_REF_STYLE, WD_STYLE_TYPE.CHARACTER)
    va = ensure_ordered_child(ensure_rpr(ref.element), 'vertAlign', _RPR_ORDER)
    va.set(W('val'), 'superscript')
    logger.info(f' ✓ {_FOOTNOTE_REF_STYLE}: 已设为上标')


def generate_template(requirements_text, output_path, margins=None, lang="zh"):
    """主流程：生成模板"""
    output_path = str(output_path)
    logger.info('─' * 50)
    logger.info('生成学术 Word 模板')
    logger.info('─' * 50)

    spec = parse_requirements(requirements_text)
    page_size = spec.page_size
    reqs, default_latin = spec.styles, spec.default_latin
    if not reqs:
        raise TemplateError(
            '未解析到任何样式要求；格式应为「样式名：属性」，'
            '如「一级标题：黑体 三号 加粗 不倾斜 左对齐 行距1.5倍 段前12磅 段后6磅 首行缩进0字符」。'
        )

    if spec.margins:
        margins = spec.margins
    if not margins:
        margins = {'top': 2.54, 'bottom': 2.54, 'left': 2.54, 'right': 2.54}

    logger.info(f'\n解析到 {len(reqs)} 个样式定义:')
    for r in reqs:
        desc = f'{r["style_name"]}'
        if 'font_name' in r:
            desc += f' {", ".join(r["font_name"])}'
        if 'font_size' in r:
            desc += f' {r["font_size"].pt}pt'
        logger.info(f' - {desc}')
    if default_latin:
        logger.info(f' 全局西文字体: {default_latin}')

    if _check_required_fields(spec):
        raise TemplateError('必填样式或必填字段缺失，详见上方逐条提示。')

    logger.info('\n加载纯净现代 Word 365 底模...')
    tmp_path = get_default_template()
    doc = Document(tmp_path)
    cleanup_temp(tmp_path)

    if spec.proofing_lang:
        apply_proofing_language(doc, spec.proofing_lang)
        logger.info(f' ✓ 校对语言: {spec.proofing_lang}')

    for req in reqs:
        style_name = req.get('style_name')
        if not style_name:
            continue
        modify_style(doc, style_name, _props_from_req(req, page_size, margins),
                     default_latin=default_latin)

    apply_table_format(doc, spec, default_latin=default_latin)
    apply_page_decorations(doc, spec)
    apply_footnote_format(doc, spec, page_size, margins, default_latin=default_latin)

    try:
        hyperlink_style = doc.styles['Hyperlink']
        hyperlink_style.font.color.rgb = RGBColor(0, 0, 0)
        hyperlink_style.font.underline = False
        rpr = hyperlink_style.element.find(W('rPr'))
        if rpr is not None:
            color_el = rpr.find(W('color'))
            if color_el is not None:
                clear_color_theme_attrs(color_el)
        logger.info(' ✓ Hyperlink: 已设为黑色（取消下划线）')
    except KeyError:
        pass

    # 清除 Caption 基样式的斜体：Pandoc/Quarto 默认 Caption 基样式带 italic，
    # Image Caption 和 Table Caption 继承它，导致 bold+italic 叠加
    try:
        caption_style = doc.styles['Caption']
        caption_style.font.italic = False
        logger.info(' ✓ Caption: 已清除斜体（影响 Image Caption / Table Caption）')
    except KeyError:
        pass

    # 清除 Abstract 样式的 keepNext/keepLines：
    # 结构化摘要分段后，各段 keepNext 串成"必须同页"的传递链（AbstractTitle → 各段 → Keywords），
    # 链高超过首页剩余空间时 Word 把整链下移，首页只剩标题。清除后摘要按普通段落流排版。
    try:
        abstract_style = doc.styles['Abstract']
        ppr = abstract_style.element.find(W('pPr'))
        if ppr is not None:
            removed = []
            for tag in ('keepNext', 'keepLines'):
                el = ppr.find(W(tag))
                if el is not None:
                    ppr.remove(el)
                    removed.append(tag)
            if removed:
                logger.info(f' ✓ Abstract: 已清除 {" / ".join(removed)}（结构化摘要分段不再触发首页分页）')
    except KeyError:
        pass

    for style_el in doc.styles._element:
        style_id = style_el.get(W('styleId'))
        if not style_id:
            continue

        is_char_style = 'Char' in style_id and any(
            style_id.startswith(p) for p in ('Title', 'Subtitle', 'Heading')
        )
        is_unused_heading = style_id in ('Subtitle', 'TOCHeading') or (
            style_id and style_id.startswith('Heading') and style_id[-1] in '456789'
        )

        if is_char_style or is_unused_heading:
            rpr = style_el.find(W('rPr'))
            if rpr is not None:
                rfonts = rpr.find(W('rFonts'))
                color_el = rpr.find(W('color'))
                had_font = rfonts is not None and clear_font_theme_attrs(rfonts)
                had_color = color_el is not None and clear_color_theme_attrs(color_el)
                if color_el is not None and color_el.get(W('val')) not in ('000000', 'auto'):
                    color_el.set(W('val'), '000000')
                    had_color = True
                if had_font and is_char_style:
                    logger.info(f' ✓ {style_id}: 已清除主题字体')
                if had_color and is_char_style:
                    logger.info(f' ✓ {style_id}: 已清除主题色')

    from docx.shared import Cm
    w_cm, h_cm = PAGE_SIZES_CM.get(page_size, PAGE_SIZES_CM['a4'])
    pg_w, pg_h = Cm(w_cm), Cm(h_cm)

    # 装订线**折进左边距**，不写 pgMar/@w:gutter。
    # Word 打印时两者等效（单节、不开镜像页边距时，"装订线"就是在装订侧多留空白）；
    # 但 pandoc 以及本技能自己算正文宽度，用的都是 `pgSz - pgMar(left/right)`，
    # **都不认 w:gutter**：实测 A4 + 2.54cm 边距 + 装订线1cm，`{width=25cm}` 被钳成
    # 15.90cm（本应是 14.92cm），而文字起点已右移 1cm，于是图片压出右边距 **0.98cm**
    # （在 Word 渲染的 PDF 上量出来的）。折进去之后，图片宽度、公式网格表、点线制表位
    # 的落点全跟着对，且都是"算得对"而不是"再补一刀"。
    eff_left = margins['left'] + (spec.gutter_cm or 0)

    for section in doc.sections:
        section.page_width = pg_w
        section.page_height = pg_h
        section.top_margin = Cm(margins['top'])
        section.bottom_margin = Cm(margins['bottom'])
        section.left_margin = Cm(eff_left)
        section.right_margin = Cm(margins['right'])
        if spec.header_distance is not None:
            section.header_distance = Cm(spec.header_distance)
        if spec.footer_distance is not None:
            section.footer_distance = Cm(spec.footer_distance)
        # 文档网格：每页 N 行 → 每行占多少 twips，由版心高度除以行数算
        if spec.doc_grid_lines:
            grid = ensure_ordered_child(section._sectPr, 'docGrid', _SECTPR_ORDER)
            grid.set(W('type'), 'lines')
            grid.set(W('linePitch'), str(doc_grid_pitch(spec, margins)))
        w_cm = round(pg_w.inches * 2.54, 1)
        h_cm = round(pg_h.inches * 2.54, 1)
        content_w = round(w_cm - eff_left - margins['right'], 1)
        decor_cm = ''
        if spec.header_distance is not None or spec.footer_distance is not None:
            decor_cm = (f'，页眉距边界 {spec.header_distance}cm 页脚距边界 {spec.footer_distance}cm')
        extra = ''
        left_note = ''
        if spec.gutter_cm:
            extra += f'，装订线 {spec.gutter_cm}cm'
            left_note = f'（{margins["left"]}+装订线{spec.gutter_cm}）'
        if spec.doc_grid_lines:
            extra += f'，文档网格 每页 {spec.doc_grid_lines} 行'
        logger.info(f' ✓ 页面: {page_size} ({w_cm}×{h_cm}cm)，边距 上{margins["top"]}cm 下{margins["bottom"]}cm 左{eff_left}cm{left_note} 右{margins["right"]}cm，内容宽 {content_w}cm{decor_cm}{extra}')

    try:
        doc.styles['Block Text']
    except KeyError:
        doc.styles.add_style('Block Text', WD_STYLE_TYPE.PARAGRAPH)
        logger.info(' ✓ Block Text: 已创建')

    try:
        doc.styles['Source Code']
    except KeyError:
        doc.styles.add_style('Source Code', WD_STYLE_TYPE.PARAGRAPH)
        logger.info(' ✓ Source Code: 已创建')

    doc.save(output_path)

    # 保存后重新打包：只做一件事——清除 styles.xml 里的全部主题字体/颜色属性。
    # settings.xml 不在这里动：compatMode 与 mathPr 由底模写好，再注入一遍就是
    # 同一个事实写两处。
    try:
        import zipfile
        with zipfile.ZipFile(output_path, 'r') as zin:
            styles_xml = zin.read('word/styles.xml')
            has_styles_with_effects = 'word/stylesWithEffects.xml' in zin.namelist()
            other_items = {item.filename: zin.read(item.filename) for item in zin.infolist()
                           if item.filename not in ('word/styles.xml', 'word/stylesWithEffects.xml')}

        cleaned_styles_xml = strip_theme_attrs_from_xml(styles_xml)

        tmp_out = output_path + ".tmp"
        with zipfile.ZipFile(tmp_out, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
            zout.writestr('word/styles.xml', cleaned_styles_xml)
            if has_styles_with_effects:
                zout.writestr('word/stylesWithEffects.xml', cleaned_styles_xml)
            for fname, fcontent in other_items.items():
                zout.writestr(fname, fcontent)
        os.replace(tmp_out, output_path)
    except Exception as e:
        # 不吞：跳过这一步，主题字体/颜色就可能盖掉显式设置，而外表看不出异常。
        # 宁可当场失败，也不要交出一个看着正常的坏模板。
        raise RuntimeError(f'模板已写出，但「styles.xml 主题属性清理」失败，产物不合格：{e}') from e

    logger.info(f'✅ 模板已生成: {output_path} (现代 Word 365 引擎，compatMode=15)')
    logger.info(t("template.styles_only", lang))


def main():
    if len(sys.argv) < 2 or sys.argv[1] in ('-h', '--help'):
        print(__doc__)
        sys.exit(0)

    requirements_file = None
    output_path = 'template.docx'

    args = sys.argv[1:]
    i = 0
    while i < len(args):
        arg = args[i]
        if arg in ('-o', '--output'):
            i += 1
            if i < len(args):
                output_path = args[i]
        elif not arg.startswith('-') and requirements_file is None:
            requirements_file = arg
        else:
            logger.error(f'错误: 不认识的参数 "{arg}"')
            logger.error('用法: python3 generate_template.py requirements.txt -o template.docx')
            sys.exit(1)
        i += 1

    if requirements_file is None:
        logger.error('错误: 未提供格式要求文件')
        logger.error('用法: python3 generate_template.py requirements.txt -o template.docx')
        sys.exit(1)

    if not os.path.exists(requirements_file):
        logger.error(f'错误: 文件不存在 "{requirements_file}"')
        sys.exit(1)

    with open(requirements_file, 'r', encoding='utf-8') as f:
        requirements_text = f.read()

    try:
        generate_template(requirements_text, output_path)
    except TemplateError as exc:
        logger.error(f'错误: {exc}')
        sys.exit(1)


if __name__ == '__main__':
    main()

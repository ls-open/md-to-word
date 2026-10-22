"""模型 → OOXML 的写出器。

验收方式是对产物本身下断言：读 `word/document.xml`，检查结构、样式与编号
样式序列与全文文本。样式序列一致，说明「哪一段用什么样式」这层判断没有跑偏；
文本一致，说明内容一个字没丢。

能力边界单独测：没实现的结构必须抛 `UnsupportedFeature`，不许静默降级成普通段落。
"""

from __future__ import annotations

import re
import shutil
import zipfile
from pathlib import Path

import pytest

from md_to_word.model.reader import read_markdown
from md_to_word.template.builder import build_template
from md_to_word.writer import UnsupportedFeature, write_docx

REPO_ROOT = Path(__file__).resolve().parents[1]
PRESET = REPO_ROOT / "presets" / "zh-academic.txt"
DEMO = REPO_ROOT / "demo" / "00_quickstart"


def _facts(path: Path) -> tuple[list[str], str]:
    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")
    styles = re.findall(r'<w:pStyle w:val="([^"]+)"', xml)
    text = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml))
    return styles, text


@pytest.fixture()
def template(tmp_path: Path) -> Path:
    out = tmp_path / "template.docx"
    build_template(PRESET.read_text(encoding="utf-8"), out)
    return out


def test_writes_title_block_and_blocks(template: Path, tmp_path: Path) -> None:
    model = read_markdown(
        '---\ntitle: "题目"\nauthor: [张三]\n---\n\n# 一 引言\n\n正文。\n\n## 1.1 小节\n\n更多正文。\n'
    )
    out = write_docx(model, template, tmp_path / "a.docx")
    styles, text = _facts(out)
    assert styles[:2] == ["Title", "Author"]
    assert "Heading1" in styles
    assert "First Paragraph" in text or "FirstParagraph" in styles


def test_first_paragraph_after_heading_uses_first_paragraph_style(template: Path, tmp_path: Path) -> None:
    model = read_markdown("# 一 引言\n\n第一段。\n\n第二段。\n")
    styles, _ = _facts(write_docx(model, template, tmp_path / "b.docx"))
    assert styles == ["Heading1", "FirstParagraph", "BodyText"]


def test_table_header_is_marked_and_cells_use_compact(template: Path, tmp_path: Path) -> None:
    model = read_markdown("| 甲 | 乙 |\n|:--|--:|\n| 1 | 2 |\n")
    out = write_docx(model, template, tmp_path / "c.docx")
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
    assert "<w:tbl>" in xml
    assert "tblHeader" in xml, "表头行要标记跨页重复"
    assert xml.count('w:val="Compact"') == 4  # 表头两格 + 数据两格


def test_code_block_keeps_line_breaks(template: Path, tmp_path: Path) -> None:
    model = read_markdown("```\nline1\nline2\n```\n")
    out = write_docx(model, template, tmp_path / "d.docx")
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
    assert "line1" in xml and "line2" in xml


@pytest.mark.parametrize(
    ("source", "fragment"),
    [
        ("见 [@smith2024]。\n", "文献"),
    ],
)
def test_unsupported_features_fail_loudly(source: str, fragment: str, template: Path, tmp_path: Path) -> None:
    model = read_markdown(source)
    with pytest.raises(UnsupportedFeature, match=fragment):
        write_docx(model, template, tmp_path / "x.docx")


def test_lists_use_real_numbering(template: Path, tmp_path: Path) -> None:
    """编号必须是域而不是字面量，否则 Word 里按退格改不了级别、增删条目不会重编号。"""
    model = read_markdown("- 甲\n- 乙\n\n1. 一\n2. 二\n")
    out = write_docx(model, template, tmp_path / "list.docx")
    package = zipfile.ZipFile(out)
    document = package.read("word/document.xml").decode("utf-8")
    numbering = package.read("word/numbering.xml").decode("utf-8")

    assert document.count("<w:numPr>") == 4
    assert "<w:abstractNum " in numbering and "<w:num " in numbering
    assert 'w:numFmt w:val="bullet"' in numbering
    assert 'w:numFmt w:val="decimal"' in numbering


def test_ordered_list_start_override(template: Path, tmp_path: Path) -> None:
    model = read_markdown("3. 第三\n4. 第四\n")
    out = write_docx(model, template, tmp_path / "start.docx")
    numbering = zipfile.ZipFile(out).read("word/numbering.xml").decode("utf-8")
    assert 'w:startOverride w:val="3"' in numbering


def test_footnotes_are_written_as_a_real_part(template: Path, tmp_path: Path) -> None:
    """脚注不是正文里的上标文字，而是独立的 footnotes.xml 部件。"""
    model = read_markdown("# 一 章节\n\n正文有脚注[^1]。\n\n[^1]: 注释文字。\n")
    out = write_docx(model, template, tmp_path / "fn.docx")
    package = zipfile.ZipFile(out)
    assert "word/footnotes.xml" in package.namelist()

    footnotes = package.read("word/footnotes.xml").decode("utf-8")
    assert 'w:type="separator"' in footnotes, "缺分隔符，脚注区的横线会显示异常"
    assert 'w:type="continuationSeparator"' in footnotes
    assert "注释文字" in footnotes

    document = package.read("word/document.xml").decode("utf-8")
    assert 'w:id="1"' in document and "footnoteReference" in document

    # 能重新打开，说明关系与内容类型都写对了
    from docx import Document as open_docx
    open_docx(str(out))


def test_toc_field_and_section_break(template: Path, tmp_path: Path) -> None:
    """目录是 Word 域，分节是挂在 pPr 里的 sectPr——都不是普通段落。"""
    model = read_markdown(
        "{{< toc 2 >}}\n\n{{< section 页码=小写罗马 起始=1 >}}\n\n# 一 引言\n\n正文。\n"
    )
    out = write_docx(model, template, tmp_path / "sec.docx")
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")

    assert "instrText" in xml and "TOC" in xml, "目录必须是域，否则页码不会更新"
    assert xml.count("<w:sectPr") == 2, "分节段落一个 + 文末一个"
    assert '<w:pgNumType w:fmt="lowerRoman" w:start="1"/>' in xml


def test_equations_become_omml_with_numbers(template: Path, tmp_path: Path) -> None:
    """公式要写成原生 OMML（Word 里可编辑），编号按出现顺序，交叉引用解析成「式N」。"""
    model = read_markdown(
        '---\ntitle: "t"\ncrossref:\n  eq-prefix: "式"\n---\n\n# 一 引言\n\n'
        "行内 $x+y$。\n\n$$E = mc^2$$ {#eq-a}\n\n由 @eq-a 可知。\n"
    )
    out = write_docx(model, template, tmp_path / "eq.docx")
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")

    assert xml.count("<m:oMath") == 2, "行内一条 + 显示一条"
    assert "（1）" in xml, "中文稿默认全角括号"
    assert "式（1）" in xml, "交叉引用与公式旁的编号用同一份格式"


def test_equation_numbering_follows_chapters(template: Path, tmp_path: Path) -> None:
    model = read_markdown(
        '---\ntitle: "t"\ncrossref:\n  chapters: true\n---\n\n'
        "# 第一章\n\n$$a$$ {#eq-a}\n\n# 第二章\n\n$$b$$ {#eq-b}\n"
    )
    model.meta.chapters = True
    out = write_docx(model, template, tmp_path / "ch.docx")
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
    assert "（1-1）" in xml and "（2-1）" in xml, "中文稿按章编号默认用短横"


def test_equation_number_format_is_a_plain_template(template: Path, tmp_path: Path) -> None:
    """括号与分隔符不枚举变体：一行模板写什么就是什么，中英文稿都照办。"""
    model = read_markdown('---\ntitle: "t"\ncrossref:\n  chapters: true\n---\n\n# A\n\n$$a$$ {#eq-a}\n')
    out = write_docx(
        model, template, tmp_path / "fmt.docx",
        equation_number_format="[第{c}章 · {n}]",
    )
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
    assert "[第1章 · 1]" in xml


def test_english_manuscript_defaults_to_half_width(template: Path, tmp_path: Path) -> None:
    """默认外观跟稿件语种走：英文稿 (1.1)，不靠命令行开关。"""
    model = read_markdown(
        '---\ntitle: "T"\ncrossref:\n  chapters: true\n---\n\n# Intro\n\n$$a$$ {#eq-a}\n\nSee @eq-a.\n'
    )
    out = write_docx(model, template, tmp_path / "en.docx")
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
    assert "(1.1)" in xml and "Equation (1.1)" in xml


def test_chapter_format_without_chapters_fails_loudly(template: Path, tmp_path: Path) -> None:
    """模板用了 {c} 而稿件没按章编号：当场报错，不产出 （0-1）。"""
    from md_to_word.writer.docx_writer import UnsupportedFeature

    model = read_markdown('---\ntitle: "t"\n---\n\n# 一\n\n$$a$$ {#eq-a}\n')
    with pytest.raises(UnsupportedFeature, match="chapters"):
        write_docx(model, template, tmp_path / "bad.docx", equation_number_format="（{c}-{n}）")


def test_unsupported_mathml_fails_loudly() -> None:
    """转换器遇到没覆盖的 MathML 结构要当场报错，不能悄悄变成空白或纯文本。

    注意：`latex2mathml` 本身对不认识的命令很宽容（原样转成文本），所以这条边界只能
    在 MathML 这一层测——那也正是我们自己的代码所在的位置。
    """
    from md_to_word.transform.mathml2omml import UnsupportedMath, mathml_to_omml

    unsupported = (
        '<math xmlns="http://www.w3.org/1998/Math/MathML">'
        "<menclose><mi>x</mi></menclose></math>"
    )
    with pytest.raises(UnsupportedMath, match="menclose"):
        mathml_to_omml(unsupported)


def test_pipeline_needs_no_external_binary(tmp_path: Path) -> None:
    """内化的判据不是"我们也能渲染"，而是"没有任何外部二进制也能渲染"。

    依赖清单里只有纯 Python 包；这条用例跑通，就说明导出链路没偷偷调 pandoc/quarto。
    """
    from md_to_word.pipeline import ExportOptions, run_export_pipeline

    for name in ("hello.md", "requirements.txt"):
        shutil.copy(DEMO / name, tmp_path / name)
    if (DEMO / "figures").is_dir():
        shutil.copytree(DEMO / "figures", tmp_path / "figures")

    result = run_export_pipeline(
        ExportOptions(input_path=tmp_path / "hello.md", output_path=tmp_path / "out.docx")
    )

    assert result.returncode == 0, result.errors
    assert (tmp_path / "out.docx").exists()
    styles, text = _facts(tmp_path / "out.docx")
    assert "Title" in styles and text


def test_citations_render_and_bibliography_is_appended(template: Path, tmp_path: Path) -> None:
    """引文由 CSL 引擎渲染，文末列表自动生成——作者一个字都不用写。"""
    from md_to_word.transform.citations import CitationEngine

    bib = tmp_path / "refs.bib"
    bib.write_text(
        "@article{smith2024,\n  author = {Smith, John},\n  title = {A Study},\n"
        "  journal = {Nature},\n  year = {2024},\n}\n",
        encoding="utf-8",
    )
    csl = REPO_ROOT / "src" / "md_to_word" / "resources" / "csl" / "chinese-gb7714-2015-numeric.csl"
    engine = CitationEngine(bib.read_text(encoding="utf-8"), csl)

    model = read_markdown('---\ntitle: "t"\n---\n\n# 一 引言\n\n见 [@smith2024]。\n')
    out = write_docx(model, template, tmp_path / "cite.docx", citations=engine)
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
    text = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml))

    assert "[1]" in text, "正文里的引用要渲染成编号"
    assert "SMITH" in text.upper(), "文末要自动生成文献列表"
    assert 'w:val="Bibliography"' in xml, "文献条目要用 Bibliography 样式"


def test_bibtex_to_csl_json_maps_core_fields() -> None:
    from md_to_word.transform.bibtex import parse_bibtex

    entries = parse_bibtex(
        "@article{key1,\n  author = {Zhang, San and Li, Si},\n  title = {T},\n"
        "  journal = {J},\n  year = {2020},\n  pages = {1--9},\n}\n"
    )
    assert entries[0]["type"] == "article-journal"
    assert entries[0]["container-title"] == "J"
    assert entries[0]["issued"] == {"date-parts": [[2020]]}
    assert [a.get("family") for a in entries[0]["author"]] == ["Zhang", "Li"]


def test_nary_keeps_operand_inside(template: Path, tmp_path: Path) -> None:
    """`\\sum_{k} \\beta_k` 的求和号必须把被运算体吃进 `m:e`。

    否则 Word 先画一个空的求和号（带占位框），下限被挤到下一行，整条公式散架。
    """
    model = read_markdown(r"$$Y = \sum_{k} \beta_k X_k$$" + "\n")
    out = write_docx(model, template, tmp_path / "a.docx")
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
    nary = xml[xml.index("<m:nary>") : xml.index("</m:nary>")]
    assert "<m:sub>" in nary and "k" in nary.split("<m:sub>")[1].split("</m:sub>")[0]
    assert "<m:e>" in nary and "<m:sSub>" in nary, "被运算体应落在 m:e 里"


def test_ascii_quotes_become_typographic(template: Path, tmp_path: Path) -> None:
    """中文稿里的 `"…"` 直引号印出来很显眼，与中文排版的习惯不符。"""
    model = read_markdown('---\ntitle: "t"\n---\n\n本文不作"无遗漏变量偏误"的断言。\n')
    out = write_docx(model, template, tmp_path / "a.docx")
    _, text = _facts(out)
    assert "“无遗漏变量偏误”" in text


def test_cjk_citation_gap_is_trimmed(template: Path, tmp_path: Path) -> None:
    """`认知负荷 @key` 里的空格是 Markdown 语法要求，不该印进 Word。"""
    from md_to_word.config import get_csl_dir
    from md_to_word.transform.citations import CitationEngine

    engine = CitationEngine(
        "@article{smith2024,\n  author={Smith, J.},\n  title={T},\n  year={2024},\n  journal={J}\n}\n",
        get_csl_dir() / "chinese-gb7714-2015-numeric.csl",
    )
    model = read_markdown('---\ntitle: "题"\n---\n\n认知负荷 @smith2024 增加。\n', citation_keys={"smith2024"})
    out = write_docx(model, template, tmp_path / "a.docx", citations=engine)
    _, text = _facts(out)
    assert "认知负荷" in text and "认知负荷 " not in text


def _grid_widths(path: Path) -> list[list[int]]:
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(path)
    return [
        [int(g.get(qn("w:w"))) for g in t._tbl.tblGrid.findall(qn("w:gridCol"))]
        for t in doc.tables
    ]


def test_separator_dash_lengths_set_column_widths(template: Path, tmp_path: Path) -> None:
    """短横线长短写下的是列宽比例；`:---` 与 `:---:` 只差对齐，不是比例。"""
    model = read_markdown(
        "| 甲 | 乙 | 丙 |\n|:--|:--:|:------|\n| 1 | 2 | 3 |\n"
    )
    out = write_docx(model, template, tmp_path / "a.docx")
    [widths] = _grid_widths(out)
    assert widths[0] == widths[1] < widths[2]
    assert sum(widths) > 0


def test_equal_dashes_leave_the_table_to_the_word_engine(template: Path, tmp_path: Path) -> None:
    """全等 = 没指定，一列都不动，由内容自己决定。"""
    from docx import Document

    model = read_markdown("| 甲 | 乙 |\n|:---|:---:|\n| 1 | 2 |\n")
    out = write_docx(model, template, tmp_path / "a.docx")
    [widths] = _grid_widths(out)
    assert len(set(widths)) == 1
    assert Document(out).tables[0].autofit is True


def test_equation_number_placement_is_configurable(template: Path, tmp_path: Path) -> None:
    """编号默认靠右；要求文件写「公式编号：紧跟」时直接跟在公式后面。"""
    model = read_markdown(r"$$y = ax + b$$ {#eq-a}" + "\n")

    right = write_docx(model, template, tmp_path / "r.docx")
    inline = write_docx(model, template, tmp_path / "i.docx", equation_number_placement="inline")

    right_xml = zipfile.ZipFile(right).read("word/document.xml").decode("utf-8")
    inline_xml = zipfile.ZipFile(inline).read("word/document.xml").decode("utf-8")
    assert 'w:val="right"' in right_xml and "(1)" in right_xml
    assert 'w:val="right"' not in inline_xml, "紧跟时不设右制表位"
    assert "(1)" in inline_xml

"""Unit and integration tests for md_to_word.pipeline."""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

import pytest

SRC_DIR = Path(__file__).resolve().parents[1] / "src"
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

from md_to_word.pipeline import ExportOptions, run_export_pipeline


def test_pipeline_export_zh_academic(tmp_path: Path):
    fixture_dir = Path(__file__).resolve().parent / "fixtures"
    for name in ("manuscript.md", "refs.bib"):
        shutil.copy(fixture_dir / name, tmp_path / name)

    output = tmp_path / "output_zh.docx"
    options = ExportOptions(
        input_path=tmp_path / "manuscript.md",
        output_path=output,
        preset="zh-academic",
    )
    result = run_export_pipeline(options)
    assert result.returncode == 0, f"Errors: {result.errors}"
    assert output.exists()
    assert output.stat().st_size > 5000


def test_pipeline_export_en_nature(tmp_path: Path):
    fixture_dir = Path(__file__).resolve().parent / "fixtures"
    for name in ("manuscript.md", "refs.bib"):
        shutil.copy(fixture_dir / name, tmp_path / name)

    output = tmp_path / "output_en.docx"
    options = ExportOptions(
        input_path=tmp_path / "manuscript.md",
        output_path=output,
        preset="en-nature",
    )
    result = run_export_pipeline(options)
    assert result.returncode == 0, f"Errors: {result.errors}"
    assert output.exists()
    assert output.stat().st_size > 5000


def test_pipeline_fails_on_invalid_manuscript(tmp_path: Path):
    bad_md = tmp_path / "bad.md"
    bad_md.write_text("""---
title: Broken
---

## 1. Wrong heading level
![](missing_image.png)
""", encoding="utf-8")

    output = tmp_path / "bad.docx"
    options = ExportOptions(
        input_path=bad_md,
        output_path=output,
        preset="zh-academic",
    )
    result = run_export_pipeline(options)
    assert result.returncode == 1
    # 稿件是英文，门禁也按英文报（`--preset zh-academic` 只决定版式，不决定消息语言）
    assert any("image file not found" in e for e in result.errors)


def _write_bib(directory: Path) -> None:
    (directory / "refs.bib").write_text(
        "@article{smith2024,\n  author={Smith, John},\n  title={A Study},\n"
        "  year={2024},\n  journal={Journal of Tests}\n}\n",
        encoding="utf-8",
    )


def _paragraph_texts(docx_path: Path) -> list[str]:
    from docx import Document

    return [p.text for p in Document(docx_path).paragraphs]


def test_default_csl_is_shared_by_both_writers(tmp_path: Path) -> None:
    """稿件没写 `csl:` 时，两条通道必须用同一个默认样式。

    否则同一份稿子换个后端就换出两种文献表——作者只会觉得"怎么又变了"。
    """
    _write_bib(tmp_path)
    (tmp_path / "m.md").write_text(
        '---\nbibliography: refs.bib\n---\n\n# 一 引言\n\n'
        "中文正文引用一次 [@smith2024] 即可。\n",
        encoding="utf-8",
    )
    native = tmp_path / "native.docx"
    result = run_export_pipeline(
        ExportOptions(input_path=tmp_path / "m.md", output_path=native, check_md=False)
    )
    assert result.returncode == 0, result.errors
    text = "\n".join(_paragraph_texts(native))
    assert "[1]" in text, "中文稿未写 csl: 时应走 GB/T 数字制"
    assert "Smith, John. (2024)" not in text, "不该走作者-年份制"


def test_default_csl_follows_language(tmp_path: Path) -> None:
    """英文稿不写 `csl:` 走 APA：同一个机制，按字种选样式。"""
    from md_to_word.config import default_csl_path

    assert default_csl_path("这是一份中文稿件。" * 20).name == "chinese-gb7714-2015-numeric.csl"
    assert default_csl_path("This is an English manuscript. " * 20).name == "apa.csl"


def test_export_reminds_to_open_the_result(tmp_path: Path) -> None:
    """门禁全绿不等于版面合格。行高列宽、公式与题注的落版只能目视核对，
    因此每次成功导出都必须附一条核对提醒，使其成为交付流程的一环。

    措辞只给一句整体提醒，不列举该看哪里：逐项列举既啰嗦，又会让「看过了」变成
    「按清单看过、清单外没看」——真要看的是整份版面。"""
    (tmp_path / "m.md").write_text('---\n---\n\n# 一 引言\n\n正文。\n', encoding="utf-8")
    result = run_export_pipeline(
        ExportOptions(input_path=tmp_path / "m.md", output_path=tmp_path / "o.docx", preset="zh-academic")
    )
    assert result.returncode == 0, result.errors
    assert any("版面效果仍须在 Word 中人工核验" in m for m in result.messages), result.messages


def test_image_paths_are_relative_to_the_manuscript(tmp_path: Path, monkeypatch) -> None:
    """稿里的图片路径相对稿件目录，不相对进程 cwd——否则换个目录调用就全是"图不存在"。"""
    work = tmp_path / "paper"
    (work / "figures").mkdir(parents=True)
    shutil.copy(
        Path(__file__).resolve().parents[1] / "demo/01_chinese_thesis/figures/qa_pipeline.png",
        work / "figures" / "x.png",
    )
    (work / "m.md").write_text(
        '---\n---\n\n![](figures/x.png){width=8cm}\n\n%% 图片题注 图 1 示意图\n',
        encoding="utf-8",
    )
    monkeypatch.chdir(tmp_path)  # 站在别的目录调用
    result = run_export_pipeline(
        ExportOptions(input_path=work / "m.md", output_path=work / "o.docx", preset="zh-academic")
    )
    assert result.returncode == 0, result.errors
    assert "drawing" in __import__("zipfile").ZipFile(work / "o.docx").read("word/document.xml").decode()


def test_bad_requirements_return_error_instead_of_exiting(tmp_path: Path) -> None:
    """作为库调用时，写错的格式要求必须返回错误码，不能 sys.exit 掉调用方的进程。"""
    (tmp_path / "m.md").write_text('---\n---\n\n# 一\n\n正文。\n', encoding="utf-8")
    (tmp_path / "req.txt").write_text("正文：宋体 小四\n", encoding="utf-8")

    try:
        result = run_export_pipeline(
            ExportOptions(input_path=tmp_path / "m.md", output_path=tmp_path / "o.docx",
                          requirements_path=tmp_path / "req.txt")
        )
    except SystemExit:  # pragma: no cover - 旧实现会走到这里
        pytest.fail("写错的 requirements 不应终止进程，应返回 returncode=1")

    assert result.returncode == 1
    assert any("格式要求" in e for e in result.errors), result.errors


def test_custom_template_must_be_single_section(tmp_path: Path) -> None:
    """第三方模板只取正文那一节的页面几何与页眉页脚；多节模板会整体错位，必须当场拒绝。"""
    from docx import Document

    multi = tmp_path / "multi.docx"
    doc = Document()
    doc.add_paragraph("第一节")
    doc.add_section()
    doc.add_paragraph("第二节")
    doc.save(str(multi))

    (tmp_path / "m.md").write_text("# 一\n\n正文。\n", encoding="utf-8")
    result = run_export_pipeline(
        ExportOptions(input_path=tmp_path / "m.md", output_path=tmp_path / "o.docx", template_path=multi)
    )
    assert result.returncode == 1
    assert any("单节" in e for e in result.errors), result.errors


def test_user_csl_direct_formatting_is_stripped(tmp_path: Path) -> None:
    """用户自带的 CSL 里若有 hanging-indent 之类的直接格式，导出时要忽略它。

    直接格式的优先级高于样式：不剥，同一份 format.txt 换个 CSL，参考文献的
    悬挂缩进和行距就跟着变——「同一份稿件重跑一致」当场作废。
    """
    (tmp_path / "refs.bib").write_text(
        "@article{a2020, author={Doe, Jane}, title={A Paper}, journal={J}, year={2020}}\n",
        encoding="utf-8",
    )
    (tmp_path / "mine.csl").write_text(
        '<?xml version="1.0" encoding="utf-8"?>\n'
        '<style xmlns="http://purl.org/net/xbiblio/csl" class="in-text" version="1.0">\n'
        "<info><title>Mine</title><id>http://example.com/mine</id></info>\n"
        "<citation><layout><text variable=\"title\"/></layout></citation>\n"
        '<bibliography hanging-indent="true" second-field-align="flush">\n'
        "<layout><text variable=\"title\"/></layout></bibliography></style>\n",
        encoding="utf-8",
    )
    (tmp_path / "m.md").write_text(
        "---\nbibliography: refs.bib\ncsl: mine.csl\n---\n\n# 一 引言\n\n见 [@a2020]。\n\n"
        "%% 参考文献标题 参考文献\n",
        encoding="utf-8",
    )

    result = run_export_pipeline(ExportOptions(input_path=tmp_path / "m.md", output_path=tmp_path / "o.docx"))
    assert result.returncode == 0, result.errors
    assert any("hanging-indent" in m for m in result.messages), result.messages


def test_default_preset_follows_the_manuscript_language(tmp_path: Path) -> None:
    """不给 -r 也不给 --preset 时，中文稿走 zh-academic、英文稿走 en-nature。

    英文稿套上中文学位论文的格式（首行缩进两字符 + 宋体 fallback）是错的，而这个错
    看不出来：文档能打开、样式都在，只是每段都被顶了两个字符。
    """
    import re
    import zipfile

    def body_text_style(docx: Path) -> str:
        with zipfile.ZipFile(docx) as archive:
            styles = archive.read("word/styles.xml").decode("utf-8")
        match = re.search(r'<w:style [^>]*w:styleId="BodyText".*?</w:style>', styles, re.S)
        assert match, "模板里没有 Body Text 样式"
        return match.group(0)

    (tmp_path / "en.md").write_text("# 1 Introduction\n\nEnglish body.\n", encoding="utf-8")
    en = run_export_pipeline(
        ExportOptions(input_path=tmp_path / "en.md", output_path=tmp_path / "en.docx")
    )
    assert en.returncode == 0, en.errors
    assert any("en-nature" in m for m in en.messages), en.messages
    assert 'w:firstLineChars="0"' in body_text_style(tmp_path / "en.docx")

    (tmp_path / "zh.md").write_text("# 一 引言\n\n中文正文。\n", encoding="utf-8")
    zh = run_export_pipeline(
        ExportOptions(input_path=tmp_path / "zh.md", output_path=tmp_path / "zh.docx")
    )
    assert zh.returncode == 0, zh.errors
    assert any("zh-academic" in m for m in zh.messages), zh.messages
    assert 'w:firstLineChars="200"' in body_text_style(tmp_path / "zh.docx")


def test_missing_output_directory_is_reported_not_raised(tmp_path: Path) -> None:
    """输出目录不存在时要给一句话，而不是等流水线跑完再抛 FileNotFoundError。

    对着 traceback 的人不会想到自己只是漏建了一个目录。
    """
    manuscript = tmp_path / "m.md"
    manuscript.write_text("# 引言\n\n正文。\n", encoding="utf-8")

    result = run_export_pipeline(
        ExportOptions(
            input_path=manuscript,
            output_path=tmp_path / "no_such_dir" / "out.docx",
            preset="zh-academic",
        )
    )

    assert result.returncode == 1
    assert any("输出目录不存在" in e for e in result.errors), result.errors
def test_template_reader_rejects_non_docx(tmp_path: Path) -> None:
    """把 .md 当成模板传进去时给一句话，而不是 BadZipFile 的 traceback。"""
    from md_to_word.template.extractor import main as extract_main

    not_docx = tmp_path / "notes.md"
    not_docx.write_text("# 引言\n", encoding="utf-8")

    assert extract_main([str(not_docx)]) == 1


def test_library_and_cli_reverse_engineer_identically(tmp_path: Path, capsys) -> None:
    """反推要求文件只有一份实现：库调用与命令行必须给出同一份草稿。

    从前 `main()` 与 `extract_spec_from_docx()` 各写了一遍，改一边另一边就悄悄落后；
    这条把它们钉在同一份输出上。
    """
    import io
    from contextlib import redirect_stdout

    from md_to_word.template import extract_spec_from_docx
    from md_to_word.template.extractor import main as extract_main

    source = tmp_path / "seed.docx"
    assert run_export_pipeline(
        ExportOptions(
            input_path=_write_manuscript(tmp_path),
            output_path=source,
            preset="zh-academic",
        )
    ).returncode == 0

    from_library = extract_spec_from_docx(source)
    buffer = io.StringIO()
    with redirect_stdout(buffer):
        assert extract_main([str(source)]) == 0
    from_cli = buffer.getvalue()

    assert from_cli == from_library


def _write_manuscript(tmp_path: Path) -> Path:
    manuscript = tmp_path / "seed.md"
    manuscript.write_text("%% 文档标题 题目\n\n# 一\n\n正文。\n", encoding="utf-8")
    return manuscript


def test_unpicked_txt_files_are_named(tmp_path: Path) -> None:
    """没传 -r 时只认 `format.txt`；旁边的 .txt 一概不猜，但要在提醒里点名。

    猜「唯一那个 .txt 大概是要求文件」看着贴心，实则是把笔记当格式用——静默用错格式，
    比直接不认它难查得多。不猜，但也不能装作没看见：用户以为它生效了才是真问题。
    """
    (tmp_path / "doc.md").write_text("# 引言\n\n正文。\n", encoding="utf-8")
    (tmp_path / "随手记.txt").write_text("这是笔记\n", encoding="utf-8")

    result = run_export_pipeline(
        ExportOptions(input_path=tmp_path / "doc.md", output_path=tmp_path / "doc.docx")
    )

    assert result.returncode == 0, result.errors
    assert any("随手记.txt" in msg for msg in result.messages), result.messages
    assert any("格式要求" in msg for msg in result.messages), result.messages


def test_contract_cheatsheet_needs_no_manuscript(capsys) -> None:
    """只装了包、手上没有技能文档的调用方，要能自己问出稿件怎么写。"""
    from md_to_word.cli import main

    assert main(["--contract"]) == 0
    out = capsys.readouterr().out
    assert "%% 样式名" in out and "文档标题" in out and "@eq-" in out


def test_json_output_is_one_machine_readable_object(tmp_path: Path, capsys) -> None:
    """Agent 与脚本读键，不读措辞：stdout 只有一行 JSON，键名固定。"""
    import json

    from md_to_word.cli import main

    fixture_dir = Path(__file__).resolve().parent / "fixtures"
    for name in ("manuscript.md", "refs.bib"):
        shutil.copy(fixture_dir / name, tmp_path / name)

    code = main([str(tmp_path / "manuscript.md"), "-o", str(tmp_path / "out.docx"), "--json"])
    payload = json.loads(capsys.readouterr().out)

    assert code == payload["returncode"] == 0
    assert payload["ok"] is True
    assert payload["output"].endswith("out.docx")
    assert payload["errors"] == []
    assert set(payload) == {"ok", "returncode", "input", "output", "pdf", "lang", "messages", "errors"}


def test_json_output_reports_a_missing_manuscript(tmp_path: Path, capsys) -> None:
    import json

    from md_to_word.cli import main

    code = main([str(tmp_path / "nope.md"), "--json"])
    payload = json.loads(capsys.readouterr().out)

    assert code == payload["returncode"] == 1
    assert payload["ok"] is False and payload["errors"]
    assert payload["output"] is None


def test_guide_is_readable_without_the_skill_directory(capsys) -> None:
    """只装了 pip 包的人也要能一条命令读到‘按什么顺序做’。"""
    from md_to_word.cli import main

    assert main(["--guide"]) == 0
    out = capsys.readouterr().out
    assert "Read this before the first run" in out and "## 1. The work, in order" in out


def test_guide_and_skill_entry_are_one_text() -> None:
    """技能用户读 SKILL.md、pip 用户读 guide.md，两份必须是同一份流程。"""
    root = Path(__file__).resolve().parents[1]
    body = (root / "SKILL.md").read_text(encoding="utf-8").split("---", 2)[2].lstrip("\n")
    assert (root / "src/md_to_word/resources/guide.md").read_text(encoding="utf-8") == body


def test_init_writes_an_exportable_pair(tmp_path: Path, capsys) -> None:
    """第一条命令就要能出成品：init 生成的稿件与格式声明必须直接导出成功。"""
    from md_to_word.cli import main

    assert main(["init", str(tmp_path)]) == 0
    assert (tmp_path / "format.txt").exists() and (tmp_path / "hello.md").exists()

    out = tmp_path / "hello.docx"
    assert main([str(tmp_path / "hello.md"), "-o", str(out)]) == 0
    assert out.exists()
    # 第一次导出不该有任何告警：样例本身必须是范本，不能一上手就「折行段落」。
    assert "⚠️" not in capsys.readouterr().err


def test_init_does_not_overwrite_and_speaks_the_chosen_language(tmp_path: Path, capsys) -> None:
    from md_to_word.cli import main

    assert main(["init", str(tmp_path), "--lang", "en"]) == 0
    assert "Created:" in capsys.readouterr().out
    assert (tmp_path / "hello.md").read_text(encoding="utf-8").startswith("%% Title")

    code = main(["init", str(tmp_path)])
    assert code == 1
    assert "未覆盖" in capsys.readouterr().out
    assert main(["init", str(tmp_path), "--force"]) == 0


def test_init_sample_follows_the_preset_not_the_message_language(tmp_path: Path, capsys) -> None:
    """格式词表是英文（Table Caption、Body Text），样例里也必须是英文样式名。"""
    from md_to_word.cli import main

    assert main(["init", str(tmp_path), "--preset", "en-nature"]) == 0
    assert (tmp_path / "hello.md").read_text(encoding="utf-8").startswith("%% Title")
    assert main([str(tmp_path / "hello.md"), "-o", str(tmp_path / "out.docx")]) == 0


def test_init_unknown_preset_is_reported(tmp_path: Path, capsys) -> None:
    from md_to_word.cli import main

    assert main(["init", str(tmp_path), "--preset", "nope"]) == 1
    assert "nope" in capsys.readouterr().out

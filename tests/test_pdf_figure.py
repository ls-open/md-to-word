"""PDF 图回归：300 dpi、无宽度时按物理尺寸、不支持的格式在预检拦下（原理见 references/design.md §1.6.1）。"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

import pypdfium2 as pdfium

from md_to_word.model.reader import read_markdown
from md_to_word.template.builder import build_template
from md_to_word.utils.pdf_figure import (
    DEFAULT_DPI,
    PdfFigureError,
    is_pdf,
    page_count,
    rasterize,
)
from md_to_word.validator.check import validate_manuscript
from md_to_word.writer import UnsupportedFeature, write_docx

REPO_ROOT = Path(__file__).resolve().parents[1]
PRESET = REPO_ROOT / "presets" / "zh-academic.txt"


def _make_pdf(path: Path, pages: int = 1, width: int = 200, height: int = 100) -> Path:
    pdf = pdfium.PdfDocument.new()
    for _ in range(pages):
        pdf.new_page(width, height)
    pdf.save(str(path))
    pdf.close()
    return path


def _png_size(path: Path) -> tuple[int, int]:
    head = path.read_bytes()[:24]
    assert head[:8] == b"\x89PNG\r\n\x1a\n"
    return int.from_bytes(head[16:20], "big"), int.from_bytes(head[20:24], "big")


@pytest.fixture()
def template(tmp_path: Path) -> Path:
    out = tmp_path / "template.docx"
    build_template(PRESET.read_text(encoding="utf-8"), out)
    return out


def test_recognises_only_the_pdf_suffix() -> None:
    assert is_pdf("fig.pdf") and is_pdf("FIG.PDF")
    assert not is_pdf("fig.png")


def test_renders_at_the_declared_width_in_300_dpi(tmp_path: Path) -> None:
    pdf = _make_pdf(tmp_path / "fig.pdf")
    png = rasterize(pdf, width_cm=10)
    try:
        width, height = _png_size(png)
        # 10cm 在 300 dpi 上就是 1181 像素，宽高比跟着页面走（200×100 pt = 2:1）
        expected = round(10 / 2.54 * DEFAULT_DPI)
        assert width == expected
        assert height == pytest.approx(expected / 2, abs=2)
    finally:
        png.unlink()


def test_without_a_width_it_keeps_the_page_size(tmp_path: Path) -> None:
    """没有声明宽度就按页面物理尺寸落版：200×100 pt = 7.06×3.53 cm。"""
    pdf = _make_pdf(tmp_path / "fig.pdf")
    png = rasterize(pdf)
    try:
        from docx import Document

        doc = Document()
        doc.add_paragraph().add_run().add_picture(str(png))
        shape = doc.inline_shapes[0]
        assert shape.width.cm == pytest.approx(7.06, abs=0.05)
        assert shape.height.cm == pytest.approx(3.53, abs=0.05)
    finally:
        png.unlink()


def test_the_same_pdf_renders_to_the_same_bytes(tmp_path: Path) -> None:
    """确定性是这个项目的前提：同一份 PDF 每次导出的字节必须一致。"""
    pdf = _make_pdf(tmp_path / "fig.pdf")
    first = rasterize(pdf, width_cm=8)
    second = rasterize(pdf, width_cm=8)
    try:
        assert first.read_bytes() == second.read_bytes()
    finally:
        first.unlink()
        second.unlink()


def test_counts_pages(tmp_path: Path) -> None:
    assert page_count(_make_pdf(tmp_path / "two.pdf", pages=2)) == 2


def test_a_broken_pdf_says_why(tmp_path: Path) -> None:
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"not a pdf at all")
    with pytest.raises(PdfFigureError):
        page_count(broken)
    with pytest.raises(PdfFigureError):
        rasterize(broken)


def test_writer_embeds_the_rendered_png(template: Path, tmp_path: Path) -> None:
    _make_pdf(tmp_path / "fig.pdf")
    (tmp_path / "m.md").write_text(
        "# 一\n\n![](fig.pdf){width=10cm}\n\n%% 图片题注 图 1 向量图\n", encoding="utf-8"
    )
    model = read_markdown((tmp_path / "m.md").read_text(encoding="utf-8"))
    out = tmp_path / "out.docx"
    write_docx(model, template, out, base_dir=tmp_path)

    media = [n for n in zipfile.ZipFile(out).namelist() if n.startswith("word/media/")]
    assert media == ["word/media/image1.png"], media


def test_writer_reports_an_unrenderable_pdf(template: Path, tmp_path: Path) -> None:
    """校验被 --no-check 关掉时，写出器的失败也要是一句人话，不是 traceback。"""
    (tmp_path / "broken.pdf").write_bytes(b"not a pdf")
    (tmp_path / "m.md").write_text("# 一\n\n![](broken.pdf)\n", encoding="utf-8")
    model = read_markdown((tmp_path / "m.md").read_text(encoding="utf-8"))
    with pytest.raises(UnsupportedFeature):
        write_docx(model, template, tmp_path / "out.docx", base_dir=tmp_path)


BASE = '---\n---\n\n# 一 引言\n\n正文。\n'


def test_gate_blocks_a_format_word_cannot_hold(tmp_path: Path) -> None:
    (tmp_path / "fig.svg").write_text("<svg/>", encoding="utf-8")
    (tmp_path / "m.md").write_text(BASE + "\n![](fig.svg)\n", encoding="utf-8")
    report = validate_manuscript(tmp_path / "m.md")
    assert not report.valid
    assert any(".svg" in error for error in report.errors)


def test_gate_blocks_an_unreadable_pdf(tmp_path: Path) -> None:
    (tmp_path / "bad.pdf").write_bytes(b"not a pdf")
    (tmp_path / "m.md").write_text(BASE + "\n![](bad.pdf)\n", encoding="utf-8")
    report = validate_manuscript(tmp_path / "m.md")
    assert not report.valid
    assert any("bad.pdf" in error for error in report.errors)


def test_gate_warns_but_passes_a_multipage_pdf(tmp_path: Path) -> None:
    """多页多半是把整本插图集存成了一个文件；取第 1 页是看得见的结果，提醒即可。"""
    _make_pdf(tmp_path / "two.pdf", pages=2)
    (tmp_path / "m.md").write_text(BASE + "\n![](two.pdf)\n", encoding="utf-8")
    report = validate_manuscript(tmp_path / "m.md")
    assert report.valid, report.errors
    assert any("2 页" in warning for warning in report.warnings)


def test_gate_passes_a_single_page_pdf(tmp_path: Path) -> None:
    _make_pdf(tmp_path / "one.pdf")
    (tmp_path / "m.md").write_text(BASE + "\n![](one.pdf){width=10cm}\n", encoding="utf-8")
    report = validate_manuscript(tmp_path / "m.md")
    assert report.valid, report.errors
    assert not report.warnings


def test_rasterize_cache_hit_and_consistency(tmp_path: Path) -> None:
    """同一张 PDF 图重复导出时命中内存缓存，且内容逐字节一致。"""
    from md_to_word.utils.pdf_figure import _png_cache, clear_cache

    clear_cache()
    pdf = _make_pdf(tmp_path / "cached.pdf", width=150, height=80)
    assert len(_png_cache) == 0

    out1 = rasterize(pdf, width_cm=5.0)
    try:
        assert len(_png_cache) == 1
        bytes1 = out1.read_bytes()
    finally:
        out1.unlink(missing_ok=True)

    # 再次调用命中缓存
    out2 = rasterize(pdf, width_cm=5.0)
    try:
        assert len(_png_cache) == 1
        bytes2 = out2.read_bytes()
        assert bytes1 == bytes2
    finally:
        out2.unlink(missing_ok=True)


def test_rasterize_cache_is_capped_by_bytes(tmp_path: Path, monkeypatch) -> None:
    """条目数与总字节双上限：超出后淘汰最早写入的条目。"""
    from md_to_word.utils import pdf_figure

    pdf_figure.clear_cache()
    monkeypatch.setattr(pdf_figure, "_CACHE_MAX_BYTES", 1)

    first = _make_pdf(tmp_path / "a.pdf", width=150, height=80)
    second = _make_pdf(tmp_path / "b.pdf", width=150, height=40)

    for pdf in (first, second):
        out = pdf_figure.rasterize(pdf, width_cm=5.0)
        out.unlink(missing_ok=True)

    assert len(pdf_figure._png_cache) == 1, "超过字节上限时最旧的一条要被换出去"
    assert pdf_figure._CACHE_BYTES == len(next(iter(pdf_figure._png_cache.values())))

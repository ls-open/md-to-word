"""tests/test_v3_boundary_fixes.py — V3 纯自研引擎边界与缺陷修复回归测试集。"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as open_docx

from md_to_word.model.reader import read_markdown
from md_to_word.pipeline import ExportOptions, run_export_pipeline
from md_to_word.transform.citations import CitationEngine
from md_to_word.transform.mathml2omml import latex_to_omml
from md_to_word.utils.skill_utils import dialect_issues, malformed_table_issues
from md_to_word.writer.docx_writer import write_docx

REPO_ROOT = Path(__file__).resolve().parents[1]
TPL_BASE = REPO_ROOT / "src" / "md_to_word" / "resources" / "pure_modern_base.docx"
CSL_GB = REPO_ROOT / "src" / "md_to_word" / "resources" / "csl" / "chinese-gb7714-2015-numeric.csl"


def test_footnote_with_high_precision_number_passes_postflight(tmp_path: Path) -> None:
    """P0 修复回归：正文脚注中的高精度数据（如 p=0.00125）不应被 postflight 误报丢失。"""
    md_content = (
        "---\n"
        "bibliography:\n"
        "---\n\n"
        "正文引用脚注[^note1]。\n\n"
        "[^note1]: 实验实测关键数值为 3.14159265，必须完整保留。\n"
    )
    md_file = tmp_path / "paper.md"
    md_file.write_text(md_content, encoding="utf-8")

    opts = ExportOptions(
        input_path=md_file,
        preset="zh-academic",
        check_md=True,
        postflight=True,
        audit=False,
    )
    res = run_export_pipeline(opts)
    assert res.returncode == 0, f"导出失败: {res.errors}"
    assert res.output_docx and res.output_docx.is_file()


def test_table_with_escaped_pipe_and_code_pipe(tmp_path: Path) -> None:
    r"""P0 修复回归：单元格中包含转义竖线 \| 或代码 `cat | grep` 不应导致写出器越界崩溃。"""
    md_content = (
        "| 命令 | 说明 |\n"
        "| :--- | :--- |\n"
        "| `cat a | grep b` | 过滤管道符 |\n"
        "| 选项 \\| 参数 | 含转义竖线 |\n"
    )
    # 门禁检查无列数不一致报错
    assert malformed_table_issues(md_content) == []

    # 写出器正常生成
    doc = read_markdown(md_content)
    out_docx = tmp_path / "table_pipe.docx"
    write_docx(doc, TPL_BASE, out_docx)
    assert out_docx.is_file()

    # 读取验证单元格内容
    doc_docx = open_docx(out_docx)
    table = doc_docx.tables[0]
    assert len(table.rows) == 3
    assert len(table.columns) == 2
    # 转义要真的被还原成竖线，而不是把反斜杠印出来
    assert table.cell(1, 0).text == "cat a | grep b"
    assert table.cell(2, 0).text == "选项 | 参数"

    # 反斜杠写在行内代码里也一样：表格层的转义在代码段里同样生效（GFM 的规矩）
    code_escaped = read_markdown("| 命令 | 说明 |\n| :--- | :--- |\n| `a \\| b` | 代码里的竖线 |\n")
    out2 = write_docx(code_escaped, TPL_BASE, tmp_path / "table_pipe2.docx")
    assert open_docx(out2).tables[0].cell(1, 0).text == "a | b"


def test_mathml_accents_and_limits_supported() -> None:
    """P1 修复回归：支持学术高频符号（向量、帽子、均值、导数及上下限大算符）。"""
    # 向量、帽子、均值、导数
    accents = [
        r"\vec{x}",
        r"\hat{y}",
        r"\bar{z}",
        r"\dot{x}",
        r"\ddot{y}",
        r"\underline{u}",
    ]
    for expr in accents:
        nodes = latex_to_omml(expr)
        assert len(nodes) >= 1, f"转换空: {expr}"

    # 带 limits 的求和
    nary_nodes = latex_to_omml(r"\sum\limits_{i=1}^{n} x_i")
    assert len(nary_nodes) >= 1


def test_mover_distinguishes_accent_from_overset() -> None:
    """重音与「上面压一个表达式」是两件事，判据是上面的节点类型，不是文字长度。

    `\\overset{def}{=}` 的上面是一个 mrow（text 为空），按长度判会被当重音符，
    把 `=` 排成 `≙`；`\\overset{*}{x}` 只写一个字符，同样不是重音。
    """
    from lxml import etree

    def kinds(latex: str) -> list[str]:
        nodes = latex_to_omml(latex)
        return [etree.QName(n).localname for n in nodes]

    assert "acc" in kinds(r"\hat{x}")
    assert "acc" in kinds(r"\bar{x}")
    assert "limUpp" in kinds(r"\overset{def}{=}")
    assert "limUpp" in kinds(r"\overset{*}{x}")
    assert "limLow" in kinds(r"\underset{i=1}{\max}")

    over = [n for n in latex_to_omml(r"\overset{def}{=}") if etree.QName(n).localname == "limUpp"][0]
    text = "".join(over.itertext())
    assert "def" in text and "=" in text


def test_bibliography_heading_placement_before_appendix(tmp_path: Path) -> None:
    """P1 修复回归：`# 参考文献` 下的文献列表必须排在 `# 附录` 之前。"""
    bib_text = """
@article{test2024,
  author = {Zhang, San},
  title = {A Great Study},
  journal = {Nature},
  year = {2024}
}
"""
    engine = CitationEngine(bib_text, CSL_GB)
    md_text = (
        "正文引用 [@test2024]。\n\n"
        "# 参考文献\n\n"
        "# 附录 A\n\n"
        "附录正文内容。\n"
    )
    doc = read_markdown(md_text, citation_keys={"test2024"})
    out_docx = tmp_path / "order.docx"
    write_docx(doc, TPL_BASE, out_docx, citations=engine)

    res_doc = open_docx(out_docx)
    paras = [p.text for p in res_doc.paragraphs if p.text.strip()]
    # 找到 参考文献、文献条目 和 附录 A 的相对位置
    ref_idx = next(i for i, p in enumerate(paras) if "参考文献" in p)
    item_idx = next(i for i, p in enumerate(paras) if "ZHANG" in p or "A Great Study" in p)
    appendix_idx = next(i for i, p in enumerate(paras) if "附录 A" in p)

    assert ref_idx < item_idx < appendix_idx, f"顺序错乱: ref={ref_idx}, item={item_idx}, appendix={appendix_idx}"


def test_unclosed_display_math_blocked_by_preflight() -> None:
    """P1 修复回归：未闭合的独占行公式 $$ 应被门禁直接拦截打回，不替用户兜底吞噬后文。"""
    bad_md = (
        "第一段正文。\n\n"
        "$$\n"
        "E = mc^2\n\n"
        "第二段正文（本应是正文，但不幸被未闭合公式吞噬）。\n"
    )
    issues = dialect_issues(bad_md)
    assert any("独占行公式 `$$` 未闭合" in msg for _, msg in issues)


def test_escaped_dollar_not_falsely_flagged() -> None:
    """P2 修复回归：转义的美元符号 \\$ 不应被误判为未闭合的公式。"""
    text = "该产品的国内标价为 \\$100 美元整。"
    issues = dialect_issues(text)
    assert issues == [], f"误报: {issues}"


def test_styled_chinese_citation_gap_trimmed(tmp_path: Path) -> None:
    """P1 修复回归：加粗或斜体中文后的引文空格应被平滑移除。"""
    bib_text = """
@article{test2024,
  author = {Zhang, San},
  title = {A Great Study},
  journal = {Nature},
  year = {2024}
}
"""
    engine = CitationEngine(bib_text, CSL_GB)
    md_text = "**重要结论** [@test2024]。\n\n%% 参考文献标题 参考文献\n"
    doc = read_markdown(md_text, citation_keys={"test2024"})
    out_docx = tmp_path / "gap.docx"
    write_docx(doc, TPL_BASE, out_docx, citations=engine)

    res_doc = open_docx(out_docx)
    p0 = res_doc.paragraphs[0].text
    assert "重要结论[1]" in p0 or "重要结论" in p0 and " [1]" not in p0


def test_footnotes_are_numbered_and_images_keep_their_alt_text(tmp_path: Path) -> None:
    """脚注区缺 `w:footnoteRef` 会正文有号、注文无号；图片的替代文字以前读进来就丢了。"""
    import re
    import shutil
    import zipfile

    from md_to_word.model.reader import read_markdown
    from md_to_word.template.builder import build_template
    from md_to_word.writer import write_docx

    repo = Path(__file__).resolve().parents[1]
    pic = repo / "demo" / "01_chinese_thesis" / "figures" / "qa_pipeline.png"
    shutil.copy(pic, tmp_path / "pic.png")
    build_template(
        (repo / "presets" / "zh-academic.txt").read_text(encoding="utf-8"), tmp_path / "tpl.docx"
    )
    md = (
        "# 一\n\n正文有脚注[^n]，还有图。\n\n![系统架构图](pic.png){width=5cm}\n\n"
        "%% 图片题注 图 1 架构\n\n[^n]: 注释内容。\n"
    )
    write_docx(read_markdown(md), tmp_path / "tpl.docx", tmp_path / "o.docx", base_dir=tmp_path)

    package = zipfile.ZipFile(tmp_path / "o.docx")
    document = package.read("word/document.xml").decode("utf-8")
    footnotes = package.read("word/footnotes.xml").decode("utf-8")
    assert 'descr="系统架构图"' in document
    assert "footnoteRef" in footnotes


def test_each_ordered_list_restarts_numbering(tmp_path: Path) -> None:
    """两个独立的有序列表共用一个 numId 时，第二个会接着 3.、4. 往下编。"""
    import re
    import zipfile

    from md_to_word.model.reader import read_markdown
    from md_to_word.template.builder import build_template
    from md_to_word.writer import write_docx

    repo = Path(__file__).resolve().parents[1]
    build_template(
        (repo / "presets" / "zh-academic.txt").read_text(encoding="utf-8"), tmp_path / "tpl.docx"
    )
    md = "# 一\n\n1. 甲\n2. 乙\n\n中间正文。\n\n1. 丙\n2. 丁\n"
    write_docx(read_markdown(md), tmp_path / "tpl.docx", tmp_path / "o.docx")
    xml = zipfile.ZipFile(tmp_path / "o.docx").read("word/document.xml").decode("utf-8")
    ids = re.findall(r'<w:numId w:val="(\d+)"/>', xml)
    assert len(set(ids)) == 2, f"两个列表应各用一份编号：{ids}"


def test_an_over_wide_image_is_clamped_to_the_text_block(tmp_path: Path) -> None:
    """手册承诺「超宽会被钳成版心宽」；不钳的话图会溢出页面被裁。"""
    import shutil

    from docx import Document
    from docx.shared import Emu

    from md_to_word.model.reader import read_markdown
    from md_to_word.template.builder import build_template
    from md_to_word.writer import write_docx

    repo = Path(__file__).resolve().parents[1]
    shutil.copy(repo / "demo/01_chinese_thesis/figures/qa_pipeline.png", tmp_path / "pic.png")
    build_template(
        (repo / "presets" / "zh-academic.txt").read_text(encoding="utf-8"), tmp_path / "tpl.docx"
    )
    md = "# 一\n\n![](pic.png){width=25cm}\n\n%% 图片题注 图 1 宽图\n"
    write_docx(read_markdown(md), tmp_path / "tpl.docx", tmp_path / "o.docx", base_dir=tmp_path)

    document = Document(tmp_path / "o.docx")
    section = document.sections[0]
    limit = Emu(section.page_width - section.left_margin - section.right_margin).cm
    assert document.inline_shapes[0].width.cm == pytest.approx(limit, abs=0.02)


def test_footnotes_keep_inline_formatting_and_links(tmp_path: Path) -> None:
    """脚注区以前只有纯文本 run：删除线、下划线、上下标、等宽、链接的 URL 全丢。"""
    import zipfile

    from md_to_word.template.builder import build_template

    repo = Path(__file__).resolve().parents[1]
    build_template(
        (repo / "presets" / "zh-academic.txt").read_text(encoding="utf-8"), tmp_path / "tpl.docx"
    )
    md = (
        "# 一\n\n正文[^1]。\n\n"
        "[^1]: 含 ~~删除~~、[下划线]{.underline}、H~2~O、x^2^、`代码`、[链接](https://e.example)。\n"
    )
    write_docx(read_markdown(md), tmp_path / "tpl.docx", tmp_path / "o.docx")
    package = zipfile.ZipFile(tmp_path / "o.docx")
    notes = package.read("word/footnotes.xml").decode("utf-8")
    for marker in ("<w:strike/>", "w:u w:val", "vertAlign", "Verbatim", "hyperlink"):
        assert marker in notes, f"{marker} 没写进脚注：{notes[-300:]}"
    rels = "".join(
        package.read(name).decode("utf-8", "ignore")
        for name in package.namelist()
        if name.endswith(".rels")
    )
    assert "https://e.example" in rels


def test_link_citations_false_writes_plain_citations(tmp_path: Path) -> None:
    """YAML 写 `link-citations: false` 时不该再建内链。"""
    import zipfile

    from md_to_word.template.builder import build_template

    repo = Path(__file__).resolve().parents[1]
    (tmp_path / "refs.bib").write_text(
        "@article{a, author={Adams, Bo}, title={A}, journal={J}, year={2019}}\n", encoding="utf-8"
    )
    build_template(
        (repo / "presets" / "en-nature.txt").read_text(encoding="utf-8"), tmp_path / "tpl.docx"
    )
    md = "---\nbibliography: refs.bib\nlink-citations: false\n---\n\n# One\n\nSee [@a].\n"
    from md_to_word.config import get_csl_dir

    engine = CitationEngine(
        (tmp_path / "refs.bib").read_text(encoding="utf-8"), get_csl_dir() / "nature.csl", lang="en"
    )
    write_docx(
        read_markdown(md, citation_keys={"a"}),
        tmp_path / "tpl.docx",
        tmp_path / "o.docx",
        citations=engine,
        base_dir=tmp_path,
    )
    document = zipfile.ZipFile(tmp_path / "o.docx").read("word/document.xml").decode("utf-8")
    # 书签本身要留（文献表靠它定位），不该有的是文内那条超链接
    assert "<w:hyperlink" not in document
